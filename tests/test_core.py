import asyncio
import json
import re
import sqlite3
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from openai import APIConnectionError, APIStatusError
from telegram.error import BadRequest

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.admin import create_admin_app
from app.bot import InternalBot, _split_message
from app.company_context import CompanyContent, load_company_content
from app.config import Settings
from app.database import AIRuntimeProfile, AIModule, Company, Database, Membership, User
from app.document_ingestion import extract_uploaded_document
from app.knowledge import load_knowledge
from app.prompts import build_system_prompt
from app.providers import (
    ModuleGenerationError, ModuleProviderResolver, RuntimeCredentialError,
)
from app.role_profiles import load_role_profiles, resolve_communication_profile
from app.telegram_renderer import (
    TELEGRAM_OUTPUT_CONTRACT, TelegramResponsePart, markdown_to_telegram_html,
    prepare_response_parts, to_telegram_plain_text,
)


def _test_settings(
    tmp_path: Path,
    users_file: Path,
    companies_file: Path,
    **overrides,
) -> Settings:
    values = {
        "telegram_bot_token": "test-token",
        "ai_provider": "openai",
        "ai_api_key": "test-key",
        "ai_model": "test-model",
        "ai_base_url": None,
        "database_path": tmp_path / "test.db",
        "users_file": users_file,
        "companies_file": companies_file,
        "role_profiles_file": Path("config/role_profiles.json"),
        "project_root": Path("."),
        "history_limit": 12,
        "max_concurrent_updates": 4,
        "knowledge_max_chars": 50000,
        "max_response_chars": 12000,
        "log_level": "INFO",
        "admin_username": "admin",
        "admin_password": "strong-password",
        "admin_session_secret": "x" * 32,
        "admin_host": "127.0.0.1",
        "admin_port": 8080,
        "admin_cookie_secure": False,
    }
    values.update(overrides)
    return Settings(**values)


def test_user_sync_history_and_clear(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps(
            [
                {
                    "id": "company-a",
                    "name": "Company A",
                    "profile_file": "companies/company-a/profile.md",
                    "instruction_file": "companies/company-a/instruction.md",
                    "knowledge_dir": "companies/company-a/knowledge",
                    "active": True,
                },
                {
                    "id": "company-b",
                    "name": "Company B",
                    "active": True,
                },
            ]
        ),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text(
        json.dumps(
            [
                {
                    "telegram_id": 42,
                    "name": "DK",
                    "active": True,
                    "memberships": [
                        {
                            "company_id": "company-a",
                            "job_title": "Owner",
                            "division": "Management",
                            "role_level": "gm",
                            "communication_profile": "executive",
                            "custom_instruction": "Jawab strategis.",
                            "default": True,
                            "active": True,
                        },
                        {
                            "company_id": "company-b",
                            "job_title": "Advisor",
                            "communication_profile": "manager",
                            "active": True,
                        },
                    ],
                }
            ]
        ),
        encoding="utf-8",
    )
    database = Database(tmp_path / "bot.db")
    database.initialize()
    assert database.sync_companies(companies_file) == 2
    assert database.sync_users(users_file) == 1
    assert database.get_user(42).name == "DK"
    membership = database.get_active_membership(42)
    assert membership.company_id == "company-a"
    assert membership.communication_profile == "executive"
    assert database.set_active_company(42, "company-b").company_name == "Company B"
    assert database.get_active_membership(42).company_id == "company-b"

    database.add_message(42, "user", "Legacy tanpa company")
    database.sync_users(users_file)
    assert database.get_history(42, 10, "company-a") == [
        {"role": "user", "content": "Legacy tanpa company"},
    ]

    database.add_message(42, "user", "Halo A", "company-a")
    database.add_message(42, "assistant", "Hai A", "company-a")
    database.add_message(42, "user", "Halo B", "company-b")
    assert database.get_history(42, 10, "company-a") == [
        {"role": "user", "content": "Legacy tanpa company"},
        {"role": "user", "content": "Halo A"},
        {"role": "assistant", "content": "Hai A"},
    ]
    assert database.get_history(42, 10, "company-b") == [
        {"role": "user", "content": "Halo B"},
    ]

    raw_users = json.loads(users_file.read_text(encoding="utf-8"))
    raw_users[0]["memberships"] = [raw_users[0]["memberships"][0]]
    users_file.write_text(json.dumps(raw_users), encoding="utf-8")
    database.sync_users(users_file)
    assert database.set_active_company(42, "company-b") is None
    database.clear_history(42, "company-a")
    assert database.get_history(42, 10, "company-a") == []
    assert database.get_history(42, 10, "company-b") == [
        {"role": "user", "content": "Halo B"},
    ]


def test_knowledge_limit_and_prompt(tmp_path):
    knowledge_dir = tmp_path / "knowledge"
    knowledge_dir.mkdir()
    (knowledge_dir / "brand.md").write_text("Brand voice: ramah", encoding="utf-8")
    knowledge = load_knowledge(knowledge_dir, 1000)
    assert "brand.md" in knowledge
    assert "Brand voice: ramah" in knowledge

    profiles_file = tmp_path / "role_profiles.json"
    profiles_file.write_text(
        json.dumps(
            {
                "default_profile": "default",
                "profiles": [
                    {
                        "id": "staff",
                        "label": "Staff",
                        "response_level": "Operasional",
                        "role_aliases": ["Creator"],
                        "focus": ["langkah"],
                        "default_structure": ["tujuan", "checklist"],
                        "avoid": ["abstrak"],
                    },
                    {
                        "id": "default",
                        "label": "Default",
                        "response_level": "Seimbang",
                        "role_aliases": [],
                        "focus": [],
                        "default_structure": [],
                        "avoid": [],
                    },
                ],
            }
        ),
        encoding="utf-8",
    )
    profiles = load_role_profiles(profiles_file)
    user = User(
        telegram_id=1,
        name="Dyna",
        role="",
        division="",
        communication_profile="",
        custom_instruction="",
        active=True,
    )
    company = Company(
        company_id="amazing-malang",
        name="Amazing Malang",
        profile_file="",
        instruction_file="",
        knowledge_dir="",
        active=True,
    )
    membership = Membership(
        telegram_id=1,
        company_id="amazing-malang",
        company_name="Amazing Malang",
        job_title="Creator",
        division="Marketing",
        role_level="staff",
        communication_profile="",
        custom_instruction="Gunakan ide konkret.",
        is_default=True,
        active=True,
    )
    profile = resolve_communication_profile(
        membership.communication_profile, membership.job_title, profiles
    )
    content = CompanyContent(
        profile="Media lifestyle Malang",
        instruction="Utamakan akurasi lokasi.",
        knowledge=knowledge,
    )
    prompt = build_system_prompt(user, membership, company, content, profile)
    assert "Dyna" in prompt
    assert "Gunakan ide konkret" in prompt
    assert "Amazing Malang" in prompt
    assert "Utamakan akurasi lokasi" in prompt
    assert "Communication profile: Staff" in prompt
    assert "tujuan → checklist" in prompt
    assert '<knowledge company_id="amazing-malang">' in prompt

    module = AIModule(
        company_id="amazing-malang",
        company_name="Amazing Malang",
        module_id="marketing",
        name="Marketing",
        description="Perencanaan kampanye",
        ai_runtime_profile_id="marketing-openai",
        active=True,
    )
    module_prompt = build_system_prompt(
        user,
        membership,
        company,
        content,
        profile,
        module,
        "Gunakan funnel awareness sampai conversion.",
    )
    assert 'module_id="marketing"' in module_prompt
    assert "Gunakan funnel awareness sampai conversion." in module_prompt


def test_explicit_profile_overrides_role_alias():
    profiles = load_role_profiles(Path("config/role_profiles.json"))
    profile = resolve_communication_profile("manager", "Owner", profiles)
    assert profile.profile_id == "manager"


@pytest.mark.parametrize("role, expected", [("gm", "executive"), ("manager", "manager"), ("staff", "staff"), ("", "default"), ("unknown", "default")])
def test_bot_profile_is_derived_only_from_membership_role(tmp_path, role, expected):
    db = Database(tmp_path / "role.db")
    db.initialize()
    db.create_company("company-a", "Company A", "admin")
    user = db.create_user_with_membership(42, "DK", "company-a", "Owner", "Legacy Division", "staff", "executive", "", "admin")
    membership = replace(db.get_membership_admin(42, "company-a"), role_level=role)
    settings = _test_settings(tmp_path, tmp_path / "users.json", tmp_path / "companies.json")
    bot = InternalBot(settings, db, SimpleNamespace())
    profile = bot._communication_profile(replace(user, communication_profile="executive"), membership)
    assert profile.profile_id == expected
    prompt = build_system_prompt(user, membership, db.get_company("company-a"), CompanyContent("", "", ""), profile)
    assert "Legacy Division" not in prompt and "Division pada perusahaan aktif" not in prompt
    assert profile.as_prompt() in prompt


@pytest.mark.parametrize("role, expected", [("gm", "executive"), ("manager", "manager"), ("staff", "staff")])
def test_simplified_user_and_membership_forms_preserve_legacy_data(tmp_path, role, expected):
    db = Database(tmp_path / "simple-admin.db")
    db.initialize()
    db.create_company("company-a", "Company A", "admin")
    db.create_company("company-b", "Company B", "admin")
    settings = _test_settings(tmp_path, tmp_path / "users.json", tmp_path / "companies.json", database_path=db.path)
    with TestClient(create_admin_app(settings, db)) as client:
        client.post("/admin/login", data={"username": "admin", "password": "strong-password"})
        page = client.get("/admin/users/new")
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', page.text).group(1)
        assert 'name="division"' not in page.text and 'name="communication_profile"' not in page.text
        response = client.post("/admin/users", data={"csrf_token": csrf, "telegram_id": "42", "name": "DK",
            "active": "1", "company_id": "company-a", "job_title": "Owner", "role_level": role,
            "division": "forged division", "communication_profile": "forged"}, follow_redirects=False)
        assert response.status_code == 303
        membership = db.get_membership_admin(42, "company-a")
        assert membership.division == "" and membership.communication_profile == expected
        page = client.get("/admin/users/42/memberships/new")
        assert 'name="role_level"' in page.text
        assert 'name="division"' not in page.text and 'name="communication_profile"' not in page.text
        response = client.post("/admin/users/42/memberships", data={"csrf_token": csrf,
            "company_id": "company-b", "job_title": "Owner", "role_level": role}, follow_redirects=False)
        assert response.status_code == 303
        other = db.get_membership_admin(42, "company-b")
        assert other.division == "" and other.communication_profile == expected
        # Simulate pre-simplification data, including an override conflicting with role.
        with db._connect() as conn:
            conn.execute("UPDATE user_company_memberships SET division = ?, communication_profile = ? WHERE telegram_id = 42 AND company_id = 'company-a'",
                         ("  Legacy Division  ", "old-custom-profile"))
        db.add_message(42, "user", "keep-history", "company-a")
        before = db.get_membership_admin(42, "company-a")
        db.initialize()
        db.initialize()
        assert db.get_membership_admin(42, "company-a") == before
        path = "/admin/users/42/memberships/company-a"
        for url in [path + "/edit", "/admin/users/42/memberships/new", "/admin/users/42/edit"]:
            html = client.get(url).text
            assert 'name="division"' not in html and 'name="communication_profile"' not in html
            assert 'Legacy Division' not in html and 'old-custom-profile' not in html
        next_role = "manager" if role != "manager" else "gm"
        next_profile = "manager" if next_role == "manager" else "executive"
        values = {"csrf_token": csrf, "job_title": "Owner", "role_level": next_role,
                  "is_default": "1", "division": "replace legacy", "communication_profile": "staff"}
        assert client.post(path, data={**values, "csrf_token": "bad"}).status_code == 403
        assert client.post(path, data={**values, "role_level": "invalid"}).status_code == 400
        assert db.get_membership_admin(42, "company-a") == before
        assert client.post(path, data=values, follow_redirects=False).status_code == 303
        after = db.get_membership_admin(42, "company-a")
        assert after.division == before.division and after.communication_profile == before.communication_profile
        assert after.role_level == next_role and after.is_default and after.active
        assert db.get_membership_admin(42, "company-b") == other
        assert db.get_history(42, 10, "company-a")[0]["content"] == "keep-history"
        bot = InternalBot(settings, db, SimpleNamespace())
        user = db.get_user(42)
        assert bot._communication_profile(user, after).profile_id == next_profile
        assert bot._communication_profile(user, other).profile_id == expected
        replies = []
        async def reply_text(text): replies.append(text)
        update = SimpleNamespace(effective_user=SimpleNamespace(id=42), effective_message=SimpleNamespace(reply_text=reply_text))
        asyncio.run(bot.whoami(update, SimpleNamespace()))
        assert "Gaya jawaban (otomatis)" in replies[0] and "Division:" not in replies[0]
        assert "old-custom-profile" not in replies[0]


def test_existing_database_gets_communication_profile_migration(tmp_path):
    database_path = tmp_path / "legacy.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            """
            CREATE TABLE users (
                telegram_id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT '',
                division TEXT NOT NULL DEFAULT '',
                custom_instruction TEXT NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1,
                updated_at TEXT NOT NULL
            )
            """
        )
    database = Database(database_path)
    database.initialize()
    with sqlite3.connect(database_path) as connection:
        columns = {
            row[1] for row in connection.execute("PRAGMA table_info(users)").fetchall()
        }
    assert "communication_profile" in columns


def test_existing_messages_get_company_scope_migration(tmp_path):
    database_path = tmp_path / "legacy.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            """
            CREATE TABLE users (
                telegram_id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT '',
                division TEXT NOT NULL DEFAULT '',
                communication_profile TEXT NOT NULL DEFAULT '',
                custom_instruction TEXT NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1,
                updated_at TEXT NOT NULL
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                telegram_id INTEGER NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE user_sessions (
                telegram_id INTEGER PRIMARY KEY,
                active_company_id TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
    database = Database(database_path)
    database.initialize()
    with sqlite3.connect(database_path) as connection:
        message_columns = {
            row[1] for row in connection.execute("PRAGMA table_info(messages)").fetchall()
        }
        session_columns = {
            row[1]
            for row in connection.execute(
                "PRAGMA table_info(user_sessions)"
            ).fetchall()
        }
    assert "company_id" in message_columns
    assert "module_id" in message_columns
    assert "active_module_id" in session_columns


def test_existing_modules_get_runtime_profile_migration(tmp_path):
    database_path = tmp_path / "legacy-modules.db"
    with sqlite3.connect(database_path) as connection:
        connection.execute(
            """
            CREATE TABLE companies (
                company_id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                profile_file TEXT NOT NULL DEFAULT '',
                instruction_file TEXT NOT NULL DEFAULT '',
                knowledge_dir TEXT NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1,
                updated_at TEXT NOT NULL
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE modules (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company_id TEXT NOT NULL,
                module_id TEXT NOT NULL,
                name TEXT NOT NULL,
                description TEXT NOT NULL DEFAULT '',
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(company_id, module_id)
            )
            """
        )
    database = Database(database_path)
    database.initialize()
    with sqlite3.connect(database_path) as connection:
        module_columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(modules)").fetchall()
        }
        profile_table = connection.execute(
            "SELECT name FROM sqlite_master WHERE type='table' "
            "AND name='ai_runtime_profiles'"
        ).fetchone()
    assert "ai_runtime_profile_id" in module_columns
    assert "backup_ai_runtime_profile_id" in module_columns
    assert profile_table is not None


def test_module_access_publish_and_history_are_scoped(tmp_path):
    database = Database(tmp_path / "modules.db")
    database.initialize()
    database.create_company("company-a", "Company A", actor="admin")
    database.create_user_with_membership(
        42,
        "DK",
        "company-a",
        "Owner",
        "Management",
        "gm",
        "executive",
        "",
        actor="admin",
    )
    profile = database.create_ai_runtime_profile(
        "",
        "Marketing OpenAI",
        "openai",
        "AI_KEY_MARKETING",
        "gpt-5.4-mini",
        "",
        actor="admin",
    )
    first = database.create_module(
        "company-a",
        "",
        "Marketing",
        "Kelola campaign",
        actor="admin",
        ai_runtime_profile_id=profile.profile_id,
    )
    second = database.create_module(
        "company-a",
        "",
        "Marketing",
        "Module kedua",
        actor="admin",
        ai_runtime_profile_id=profile.profile_id,
    )
    assert first.module_id == "marketing"
    assert second.module_id == "marketing-2"

    database.set_membership_module_access(
        42, "company-a", [first.module_id], actor="admin"
    )
    assert database.list_accessible_modules(42, "company-a") == []
    database.save_module_playbook_draft(
        "company-a", first.module_id, "", actor="admin"
    )
    try:
        database.publish_module_playbook(
            "company-a", first.module_id, actor="admin"
        )
        assert False, "Playbook kosong seharusnya tidak dapat dipublikasikan"
    except ValueError as exc:
        assert "tidak boleh kosong" in str(exc)
    database.save_module_playbook_draft(
        "company-a", first.module_id, "Playbook marketing v1", actor="admin"
    )
    assert database.list_accessible_modules(42, "company-a") == []
    assert database.publish_module_playbook(
        "company-a", first.module_id, actor="admin"
    ) == 1
    accessible = database.list_accessible_modules(42, "company-a")
    assert [module.module_id for module in accessible] == ["marketing"]
    assert database.set_active_module(42, "marketing") == first

    replies: list[str] = []

    async def reply_text(value: str, **_kwargs) -> None:
        replies.append(value)

    settings = _test_settings(
        tmp_path,
        tmp_path / "users.json",
        tmp_path / "companies.json",
        database_path=tmp_path / "modules.db",
    )
    bot = InternalBot(settings, database, SimpleNamespace())
    update = SimpleNamespace(
        effective_user=SimpleNamespace(id=42),
        effective_message=SimpleNamespace(reply_text=reply_text),
    )
    asyncio.run(bot.module(update, SimpleNamespace(args=["general"])))
    assert database.get_active_module(42, "company-a") is None
    assert "General" in replies[-1]
    asyncio.run(bot.module(update, SimpleNamespace(args=["marketing"])))
    assert database.get_active_module(42, "company-a") == first
    assert "Marketing" in replies[-1]

    database.add_message(42, "user", "General", "company-a", "")
    database.add_message(42, "user", "Marketing", "company-a", "marketing")
    assert database.get_history(42, 10, "company-a", "") == [
        {"role": "user", "content": "General"}
    ]
    assert database.get_history(42, 10, "company-a", "marketing") == [
        {"role": "user", "content": "Marketing"}
    ]

    database.save_module_playbook_draft(
        "company-a", "marketing", "Playbook marketing v2", actor="admin"
    )
    assert database.get_published_module_playbook(
        "company-a", "marketing"
    ) == "Playbook marketing v1"
    assert database.publish_module_playbook(
        "company-a", "marketing", actor="admin"
    ) == 2
    versions = database.list_module_playbook_versions("company-a", "marketing")
    database.restore_module_playbook_version_to_draft(
        "company-a", "marketing", int(versions[1]["id"]), actor="admin"
    )
    state = database.get_module_playbook_admin("company-a", "marketing")
    assert state is not None
    assert state["draft_content"] == "Playbook marketing v1"
    assert state["published_content"] == "Playbook marketing v2"

    database.set_membership_module_access(42, "company-a", [], actor="admin")
    assert database.get_active_module(42, "company-a") is None
    assert database.set_active_module(42, "marketing") is None


def _company_migration_fixture(tmp_path):
    database = Database(tmp_path / "rename.db")
    database.initialize()
    for key, name in [("amazing-malang", "Amazing Malang"), ("malang-strudel", "Malang Strudel"), ("other", "Other")]:
        database.create_company(key, name, actor="admin")
    database.create_user_with_membership(42, "DK", "malang-strudel", "Owner", "Management", "gm", "executive", "Keep custom instruction", actor="admin")
    database.create_membership(42, "amazing-malang", "Owner", "Management", "gm", "executive", "", False, actor="admin")
    profile = database.create_ai_runtime_profile("", "GPT", "openai", "TEST_AI_KEY", "test-model", "", actor="admin")
    database.create_module("malang-strudel", "threads-generator", "Threads generator", "description", actor="admin", ai_runtime_profile_id=profile.profile_id)
    database.save_module_playbook_draft("malang-strudel", "threads-generator", "Playbook v1", actor="admin")
    database.publish_module_playbook("malang-strudel", "threads-generator", actor="admin")
    database.set_membership_module_access(42, "malang-strudel", ["threads-generator"], actor="admin")
    database.set_active_company(42, "malang-strudel")
    database.set_active_module(42, "threads-generator")
    database.save_company_instruction_draft("malang-strudel", "Instruction", actor="admin")
    database.publish_company_instruction("malang-strudel", actor="admin")
    database.create_knowledge_document("malang-strudel", "playbook", "Knowledge", "Knowledge content", actor="admin")
    database.publish_knowledge_document("malang-strudel", "playbook", actor="admin")
    for company, module, text in [("malang-strudel", "threads-generator", "Threads history"), ("malang-strudel", "", "General history"), ("amazing-malang", "", "Amazing history"), ("other", "", "Other history")]:
        database.add_message(42, "user", text, company, module)
    return database


def _database_rows(path):
    with sqlite3.connect(path) as connection:
        connection.row_factory = sqlite3.Row
        tables = [row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")]
        return {table: [dict(row) for row in connection.execute(f'SELECT * FROM "{table}"')] for table in tables}


def test_company_id_migration_preserves_all_data_backup_and_runtime(tmp_path):
    database = _company_migration_fixture(tmp_path)
    before = _database_rows(database.path)
    mapping = {"amazing-malang": "amz", "malang-strudel": "ms"}
    backup = Path(database.migrate_company_ids(mapping, actor="admin"))
    assert backup.stat().st_mode & 0o777 == 0o600
    assert _database_rows(backup) == before
    after = _database_rows(database.path)
    for table, rows in before.items():
        if table == "sqlite_sequence":
            continue  # Only the new audit event advances its sequence.
        expected = []
        for row in rows:
            row = dict(row)
            for column in ("company_id", "active_company_id"):
                if column in row:
                    row[column] = mapping.get(row[column], row[column])
            expected.append(row)
        assert after[table][:len(rows)] == expected
        if table != "admin_audit_events":
            assert len(after[table]) == len(rows)
    assert len(after["admin_audit_events"]) == len(before["admin_audit_events"]) + 1
    event = after["admin_audit_events"][-1]
    assert event["action"] == "company.ids_migrated"
    assert json.loads(event["details_json"])["renames"] == mapping
    assert database.get_active_membership(42).company_id == "ms"
    assert database.get_active_module(42, "ms").module_id == "threads-generator"
    assert database.get_membership_admin(42, "ms").is_default
    assert database.get_company("amz").name == "Amazing Malang"
    assert database.get_company("ms").instruction_file == "companies/malang-strudel/instruction.md"
    assert database.get_company("malang-strudel") is None
    assert database.get_history(42, 20, "ms", "threads-generator")[0]["content"] == "Threads history"
    assert database.get_history(42, 20, "ms")[0]["content"] == "General history"
    assert database.get_history(42, 20, "amz")[0]["content"] == "Amazing history"
    assert database.get_published_company_instruction("ms") == "Instruction"
    assert database.get_published_module_playbook("ms", "threads-generator") == "Playbook v1"
    assert database.get_knowledge_document_admin("ms", "playbook")["published_version_id"] is not None
    assert database.migrate_company_ids(mapping, actor="admin") is None
    assert _database_rows(database.path) == after
    assert len(list(backup.parent.glob("*.db"))) == 1
    database.initialize()
    assert database.get_active_module(42, "ms").module_id == "threads-generator"


@pytest.mark.parametrize("mapping", [
    {}, [], {"amazing-malang": "bad/id"}, {"amazing-malang": "amazing-malang"},
    {"amazing-malang": "ms", "malang-strudel": "ms"},
    {"amazing-malang": "malang-strudel", "malang-strudel": "ms"},
    {"missing": "ms"}, {"amazing-malang": "other"}, {"amazing-malang": 123},
    {"amazing-malang": "amz", "missing": "ms"},
])
def test_company_id_migration_invalid_mapping_changes_nothing(tmp_path, mapping):
    database = _company_migration_fixture(tmp_path)
    before = _database_rows(database.path)
    with pytest.raises(ValueError):
        database.migrate_company_ids(mapping, actor="admin")
    assert _database_rows(database.path) == before
    assert not (tmp_path / "backups").exists()


@pytest.mark.parametrize("failure", ["new_reference", "broken_fk", "trigger_failure", "trigger_mutation", "backup_failure"])
def test_company_id_migration_fail_closed_and_rolls_back(tmp_path, monkeypatch, failure):
    database = _company_migration_fixture(tmp_path)
    with sqlite3.connect(database.path) as connection:
        if failure == "new_reference":
            connection.execute("CREATE TABLE future_table (company_id TEXT REFERENCES companies(company_id))")
        elif failure == "broken_fk":
            connection.execute("UPDATE user_sessions SET active_company_id='missing'")
        elif failure == "trigger_failure":
            connection.execute("CREATE TRIGGER fail_rename BEFORE UPDATE ON modules BEGIN SELECT RAISE(ABORT, 'stop'); END")
        elif failure == "trigger_mutation":
            connection.execute("CREATE TRIGGER change_name AFTER UPDATE ON companies BEGIN UPDATE companies SET name='unexpected' WHERE company_id=NEW.company_id; END")
    if failure == "backup_failure":
        import app.database as database_module
        monkeypatch.setattr(database_module.tempfile, "mkstemp", lambda **kwargs: (_ for _ in ()).throw(OSError("disk full")))
    before = _database_rows(database.path)
    with pytest.raises((ValueError, sqlite3.IntegrityError, OSError)):
        database.migrate_company_ids({"amazing-malang": "amz", "malang-strudel": "ms"}, actor="admin")
    assert _database_rows(database.path) == before


def test_company_id_migration_requires_matching_audit_and_no_stale_references(tmp_path):
    database = _company_migration_fixture(tmp_path)
    mapping = {"amazing-malang": "amz", "malang-strudel": "ms"}
    database.migrate_company_ids(mapping, actor="admin")
    database.add_message(42, "user", "stale writer", "malang-strudel")
    with pytest.raises(ValueError, match="Referensi ID lama"):
        database.migrate_company_ids(mapping, actor="admin")
    with sqlite3.connect(database.path) as connection:
        connection.execute("DELETE FROM messages WHERE company_id='malang-strudel'")
        connection.execute("DELETE FROM admin_audit_events WHERE action='company.ids_migrated'")
    with pytest.raises(ValueError, match="tanpa audit"):
        database.migrate_company_ids(mapping, actor="admin")


def test_company_id_migration_rejects_orphan_target_history(tmp_path):
    database = _company_migration_fixture(tmp_path)
    database.add_message(42, "user", "unrelated old scope", "ms")
    before = _database_rows(database.path)
    with pytest.raises(ValueError, match="tidak boleh menggabungkan"):
        database.migrate_company_ids({"malang-strudel": "ms"}, actor="admin")
    assert _database_rows(database.path) == before


def test_company_id_migration_audit_failure_rolls_back(tmp_path, monkeypatch):
    import app.database as database_module
    database = _company_migration_fixture(tmp_path)
    before = _database_rows(database.path)
    def fail_audit(*args, **kwargs):
        raise sqlite3.OperationalError("audit unavailable")
    monkeypatch.setattr(database_module, "_write_audit", fail_audit)
    with pytest.raises(sqlite3.OperationalError):
        database.migrate_company_ids({"malang-strudel": "ms"}, actor="admin")
    assert _database_rows(database.path) == before
    backup = next((tmp_path / "backups").glob("*.db"))
    assert _database_rows(backup) == before


@pytest.mark.parametrize("migration", ["", '{"old":"new"}', 'invalid json'])
def test_company_id_migration_startup_before_all_runtime_writers(tmp_path, monkeypatch, migration):
    import app.main as main_module
    events = []
    settings = _test_settings(tmp_path, tmp_path / "users.json", tmp_path / "companies.json")
    monkeypatch.setattr(main_module, "load_settings", lambda: settings)
    monkeypatch.setenv("COMPANY_ID_MIGRATION", migration)
    class FakeDatabase:
        def __init__(self, path):
            pass
        def initialize(self):
            events.append("initialize")
        def migrate_company_ids(self, mapping, actor):
            assert mapping == {"old": "new"}
            events.append("migrate")
            return "verified-backup"
        def bootstrap_companies(self, path):
            events.append("bootstrap_companies")
            return 0
        def bootstrap_users(self, path):
            events.append("bootstrap_users")
            return 0
    monkeypatch.setattr(main_module, "Database", FakeDatabase)
    monkeypatch.setattr(main_module, "start_admin_server", lambda *args: events.append("admin"))
    monkeypatch.setattr(main_module, "create_provider", lambda *args: SimpleNamespace())
    application = SimpleNamespace(run_polling=lambda **kwargs: events.append("polling"))
    monkeypatch.setattr(main_module, "InternalBot", lambda *args: SimpleNamespace(build_application=lambda: application))
    if migration == "invalid json":
        with pytest.raises(json.JSONDecodeError):
            main_module.main()
        assert events == ["initialize"]
    else:
        main_module.main()
        assert events == ["initialize"] + (["migrate"] if migration else []) + ["bootstrap_companies", "bootstrap_users", "admin", "polling"]


def test_module_provider_resolver_uses_named_environment_without_fallback(
    monkeypatch,
):
    calls: list[tuple[str, str, str, str | None]] = []
    sentinel = SimpleNamespace(name="module-provider")

    def factory(provider: str, api_key: str, model: str, base_url: str | None):
        calls.append((provider, api_key, model, base_url))
        return sentinel

    profile = AIRuntimeProfile(
        profile_id="marketing-openai",
        label="Marketing OpenAI",
        provider="openai",
        api_key_env="AI_KEY_MARKETING",
        model="gpt-5.4-mini",
        base_url="",
        active=True,
    )
    resolver = ModuleProviderResolver(factory)

    monkeypatch.delenv("AI_KEY_MARKETING", raising=False)
    try:
        resolver.resolve(profile)
        assert False, "Resolver tidak boleh memakai credential global sebagai fallback"
    except RuntimeCredentialError as exc:
        assert "AI_KEY_MARKETING" in str(exc)
    assert calls == []

    monkeypatch.setenv("AI_KEY_MARKETING", "module-secret")
    assert resolver.resolve(profile) is sentinel
    assert resolver.resolve(profile) is sentinel
    assert calls == [("openai", "module-secret", "gpt-5.4-mini", None)]


@pytest.mark.parametrize("long_answer", [False, True])
def test_bot_routes_module_chat_to_module_provider_and_general_to_global(tmp_path, long_answer):
    database = Database(tmp_path / "routing.db")
    database.initialize()
    database.create_company("company-a", "Company A", actor="admin")
    database.create_user_with_membership(
        42,
        "DK",
        "company-a",
        "Owner",
        "Management",
        "gm",
        "executive",
        "",
        actor="admin",
    )
    profile = database.create_ai_runtime_profile(
        "",
        "Marketing OpenAI",
        "openai",
        "AI_KEY_MARKETING",
        "gpt-5.4-mini",
        "",
        actor="admin",
    )
    module = database.create_module(
        "company-a",
        "",
        "Marketing",
        "Campaign",
        actor="admin",
        ai_runtime_profile_id=profile.profile_id,
    )
    database.save_module_playbook_draft(
        "company-a", module.module_id, "Playbook marketing", actor="admin"
    )
    database.publish_module_playbook(
        "company-a", module.module_id, actor="admin"
    )
    database.set_membership_module_access(
        42, "company-a", [module.module_id], actor="admin"
    )
    database.set_active_module(42, module.module_id)

    class FakeProvider:
        def __init__(self, answer: str):
            self.answer = answer
            self.calls = 0

        async def generate(self, _system_prompt, _history, _user_text):
            self.calls += 1
            assert _system_prompt.endswith(TELEGRAM_OUTPUT_CONTRACT)
            return self.answer

    module_answer = "module-answer" + ("a" * 4000 if long_answer else "")
    global_provider = FakeProvider("> **global-answer**")
    module_provider = FakeProvider(
        f"**Pilihan 1**\n[[COPY_TEXT]]\n> **{module_answer}**\n"
        "[[/COPY_TEXT]]\n**Silakan pilih**"
    )
    resolver_profiles: list[str] = []

    class FakeResolver(ModuleProviderResolver):
        def resolve(self, selected_profile):
            resolver_profiles.append(selected_profile.profile_id)
            return module_provider

    settings = _test_settings(
        tmp_path,
        tmp_path / "users.json",
        tmp_path / "companies.json",
        database_path=tmp_path / "routing.db",
        project_root=tmp_path,
    )
    bot = InternalBot(settings, database, global_provider, FakeResolver())
    replies: list[str] = []
    reply_modes = []
    reject_html_once = False

    async def reply_text(value: str, **kwargs) -> None:
        nonlocal reject_html_once
        assert "parse_mode" in kwargs
        assert not kwargs.get("entities")
        assert kwargs["disable_web_page_preview"] is True
        assert len(value) <= 3600
        if reject_html_once and kwargs["parse_mode"] == "HTML":
            reject_html_once = False
            raise BadRequest("Mock HTML parse rejection")
        replies.append(value)
        reply_modes.append(kwargs["parse_mode"])

    async def send_chat_action(*_args, **_kwargs) -> None:
        return None

    message = SimpleNamespace(text="Buat campaign", reply_text=reply_text)
    update = SimpleNamespace(
        effective_user=SimpleNamespace(id=42),
        effective_message=message,
        effective_chat=SimpleNamespace(id=99),
    )
    context = SimpleNamespace(
        bot=SimpleNamespace(send_chat_action=send_chat_action)
    )
    asyncio.run(bot.chat(update, context))
    assert module_provider.calls == 1
    assert global_provider.calls == 0
    assert resolver_profiles == ["marketing-openai"]
    assert replies[0] == "<b>Pilihan 1</b>"
    assert "".join(replies[1:-1]) == module_answer
    assert replies[-1] == "<b>Silakan pilih</b>"
    assert reply_modes == ["HTML", *([None] * (len(replies) - 2)), "HTML"]
    assert database.get_history(42, 12, "company-a", module.module_id)[-1] == {
        "role": "assistant", "content": f"**Pilihan 1**\n\n{module_answer}\n\n**Silakan pilih**",
    }

    database.clear_active_module(42)
    message.text = "Pertanyaan general"
    asyncio.run(bot.chat(update, context))
    assert module_provider.calls == 1
    assert global_provider.calls == 1
    assert replies[-1] == "<blockquote><b>global-answer</b></blockquote>"
    assert reply_modes[-1] == "HTML"
    assert database.get_history(42, 12, "company-a", "")[-1] == {
        "role": "assistant", "content": "> **global-answer**",
    }

    # General can also produce a copy-ready caption, without a module-name rule.
    global_provider.answer = "[[COPY_TEXT]]\nCaption Instagram #Piko\n[[/COPY_TEXT]]"
    asyncio.run(bot.chat(update, context))
    assert replies[-1] == "Caption Instagram #Piko" and reply_modes[-1] is None

    # Even inside the same module, greetings and explanations stay formatted.
    database.set_active_module(42, module.module_id)
    module_provider.answer = "**Halo**, pilih akun `b`."
    asyncio.run(bot.chat(update, context))
    assert replies[-1] == "<b>Halo</b>, pilih akun <code>b</code>."
    assert reply_modes[-1] == "HTML"

    # A parsing rejection affects only that ordinary chunk, not the copy block
    # or the formatted guidance following it.
    reject_html_once = True
    module_provider.answer = "**Judul**\n[[COPY_TEXT]]\nCaption #Piko\n[[/COPY_TEXT]]\n**Tips**"
    asyncio.run(bot.chat(update, context))
    assert replies[-3:] == ["**Judul**", "Caption #Piko", "<b>Tips</b>"]
    assert reply_modes[-3:] == [None, None, "HTML"]


def test_company_content_is_scoped_to_configured_paths(tmp_path):
    content_dir = tmp_path / "companies" / "company-a"
    knowledge_dir = content_dir / "knowledge"
    knowledge_dir.mkdir(parents=True)
    (content_dir / "profile.md").write_text("Profil A", encoding="utf-8")
    (content_dir / "instruction.md").write_text("Instruksi A", encoding="utf-8")
    (knowledge_dir / "facts.md").write_text("Fakta A", encoding="utf-8")
    company = Company(
        company_id="company-a",
        name="Company A",
        profile_file="companies/company-a/profile.md",
        instruction_file="companies/company-a/instruction.md",
        knowledge_dir="companies/company-a/knowledge",
        active=True,
    )
    content = load_company_content(company, tmp_path, 1000)
    assert content.profile == "Profil A"
    assert content.instruction == "Instruksi A"
    assert "Fakta A" in content.knowledge


def test_admin_requires_login_and_renders_database(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Company A"}]),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text(
        json.dumps(
            [
                {
                    "telegram_id": 42,
                    "name": "DK",
                    "memberships": [
                        {
                            "company_id": "company-a",
                            "job_title": "Owner",
                            "default": True,
                        }
                    ],
                }
            ]
        ),
        encoding="utf-8",
    )
    database = Database(tmp_path / "admin.db")
    database.initialize()
    database.sync_companies(companies_file)
    database.sync_users(users_file)
    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "admin.db",
    )
    app = create_admin_app(settings, database)
    with TestClient(app) as client:
        response = client.get("/admin", follow_redirects=False)
        assert response.status_code == 303
        assert response.headers["location"] == "/admin/login"

        response = client.post(
            "/admin/login",
            data={"username": "admin", "password": "wrong-password"},
        )
        assert response.status_code == 401

        response = client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        )
        assert response.status_code == 303

        response = client.get("/admin")
        assert response.status_code == 200
        assert "Company A" in response.text
        assert "Database transition mode" in response.text
        assert 'href="/admin/static/admin.css"' in response.text

        response = client.get("/admin/static/admin.css")
        assert response.status_code == 200
        assert "--sidebar" in response.text

        response = client.get("/admin/users")
        assert response.status_code == 200
        assert "DK" in response.text
        assert "Owner" in response.text

        response = client.get("/admin/companies")
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', response.text)
        assert csrf is not None
        response = client.post(
            "/admin/companies/company-a/status",
            data={"csrf_token": csrf.group(1), "active": "0"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert "membership+aktif" in response.headers["location"]
        assert database.get_company_admin("company-a").active is True


def test_admin_company_management_and_csrf(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Company A"}]),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text("[]", encoding="utf-8")
    database = Database(tmp_path / "company-admin.db")
    database.initialize()
    assert database.bootstrap_companies(companies_file) == 1
    assert database.bootstrap_companies(companies_file) == 0

    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "company-admin.db",
    )
    app = create_admin_app(settings, database)
    with TestClient(app) as client:
        response = client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        )
        assert response.status_code == 303

        response = client.get("/admin/companies/new")
        assert response.status_code == 200
        assert 'name="company_id"' not in response.text
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', response.text)
        assert csrf is not None
        csrf_token = csrf.group(1)

        response = client.post(
            "/admin/companies",
            data={"company_id": "company-b", "name": "Company B", "active": "1"},
        )
        assert response.status_code == 403

        response = client.post(
            "/admin/companies",
            data={
                "csrf_token": csrf_token,
                "company_id": "id-yang-dipalsukan",
                "name": "Company B",
                "active": "1",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        company = database.get_company_admin("company-b")
        assert company is not None
        assert company.profile_file == "companies/company-b/profile.md"
        assert company.instruction_file == "companies/company-b/instruction.md"
        assert company.knowledge_dir == "companies/company-b/knowledge"

        response = client.post(
            "/admin/companies/company-b",
            data={"csrf_token": csrf_token, "name": "Company B Updated"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_company_admin("company-b").name == "Company B Updated"

        response = client.post(
            "/admin/companies/company-b/status",
            data={"csrf_token": csrf_token, "active": "0"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_company_admin("company-b").active is False

    actions = [row["action"] for row in database.list_admin_audit_events()]
    assert actions == [
        "company.deactivated",
        "company.updated",
        "company.created",
    ]

    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Config Name Changed"}]),
        encoding="utf-8",
    )
    assert database.bootstrap_companies(companies_file) == 0
    assert database.get_company_admin("company-a").name == "Company A"


def test_server_generated_identifiers_and_collision_suffixes(tmp_path):
    database = Database(tmp_path / "generated-identifiers.db")
    database.initialize()

    first_company = database.create_company(
        "", "Amazing Malang", actor="admin"
    )
    second_company = database.create_company(
        "", "Amazing Malang", actor="admin"
    )
    assert first_company.company_id == "amazing-malang"
    assert second_company.company_id == "amazing-malang-2"

    first_document = database.create_knowledge_document(
        first_company.company_id,
        "",
        "Target Omzet 2026",
        "Target pertama.",
        actor="admin",
    )
    second_document = database.create_knowledge_document(
        first_company.company_id,
        "",
        "Target Omzet 2026",
        "Target kedua.",
        actor="admin",
    )
    assert first_document["document_key"] == "target-omzet-2026"
    assert second_document["document_key"] == "target-omzet-2026-2"


def test_admin_activity_is_read_only_filterable_and_content_safe(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text("[]", encoding="utf-8")
    users_file = tmp_path / "users.json"
    users_file.write_text("[]", encoding="utf-8")
    database = Database(tmp_path / "activity.db")
    database.initialize()
    database.create_company("company-a", "Company A", actor="admin")
    database.create_knowledge_document(
        "company-a",
        "private-plan",
        "Private Plan",
        "RAHASIA-ISI-DOKUMEN",
        actor="admin",
    )
    database.publish_knowledge_document(
        "company-a", "private-plan", actor="admin"
    )
    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "activity.db",
    )
    app = create_admin_app(settings, database)

    with TestClient(app) as client:
        assert client.get("/admin/activity", follow_redirects=False).status_code == 303
        assert client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        ).status_code == 303
        response = client.get("/admin/activity")
        assert response.status_code == 200
        assert "Audit aktivitas" in response.text
        assert "Knowledge dipublikasikan" in response.text
        assert "company-a:private-plan" in response.text
        assert "RAHASIA-ISI-DOKUMEN" not in response.text
        assert "WIB" in response.text

        company_only = client.get("/admin/activity?category=company")
        assert company_only.status_code == 200
        assert "Company dibuat" in company_only.text
        assert "Knowledge dipublikasikan" not in company_only.text


def test_admin_user_and_membership_management(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps(
            [
                {"id": "company-a", "name": "Company A"},
                {"id": "company-b", "name": "Company B"},
            ]
        ),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text("[]", encoding="utf-8")
    database = Database(tmp_path / "user-admin.db")
    database.initialize()
    database.bootstrap_companies(companies_file)
    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "user-admin.db",
    )
    app = create_admin_app(settings, database)
    with TestClient(app) as client:
        response = client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        response = client.get("/admin/users/new")
        assert response.status_code == 200
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', response.text)
        assert csrf is not None
        csrf_token = csrf.group(1)

        user_payload = {
            "telegram_id": "123456789",
            "name": "Dyna",
            "active": "1",
            "company_id": "company-a",
            "job_title": "Content Creator",
            "division": "Marketing",
            "role_level": "staff",
            "communication_profile": "staff",
            "custom_instruction": "Jawab praktis.",
        }
        response = client.post("/admin/users", data=user_payload)
        assert response.status_code == 403
        response = client.post(
            "/admin/users",
            data={"csrf_token": csrf_token, **user_payload},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_user(123456789).name == "Dyna"
        membership_a = database.get_membership_admin(123456789, "company-a")
        assert membership_a is not None
        assert membership_a.is_default is True
        assert membership_a.communication_profile == "staff"

        response = client.get("/admin/users/123456789/edit")
        assert response.status_code == 200
        assert "Company A" in response.text
        response = client.get("/admin/users/123456789/memberships/new")
        assert response.status_code == 200
        assert "Company B" in response.text

        membership_b_payload = {
            "csrf_token": csrf_token,
            "company_id": "company-b",
            "job_title": "Project Lead",
            "division": "Growth",
            "role_level": "manager",
            "communication_profile": "manager",
            "custom_instruction": "Fokus KPI.",
            "is_default": "1",
        }
        response = client.post(
            "/admin/users/123456789/memberships",
            data=membership_b_payload,
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_membership_admin(
            123456789, "company-b"
        ).is_default is True
        assert database.get_membership_admin(
            123456789, "company-a"
        ).is_default is False
        response = client.get(
            "/admin/users/123456789/memberships/company-b/edit"
        )
        assert response.status_code == 200
        assert "Project Lead" in response.text

        response = client.post(
            "/admin/users/123456789/memberships/company-b/status",
            data={"csrf_token": csrf_token, "active": "0"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert "membership+lain" in response.headers["location"]
        assert database.get_membership_admin(123456789, "company-b").active is True

        response = client.post(
            "/admin/users/123456789/memberships/company-a",
            data={
                "csrf_token": csrf_token,
                "job_title": "Content Creator",
                "division": "Marketing",
                "role_level": "staff",
                "communication_profile": "staff",
                "custom_instruction": "Jawab praktis.",
                "is_default": "1",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_membership_admin(
            123456789, "company-a"
        ).is_default is True

        response = client.post(
            "/admin/users/123456789/memberships/company-b/status",
            data={"csrf_token": csrf_token, "active": "0"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_membership_admin(123456789, "company-b").active is False

        response = client.post(
            "/admin/users/123456789",
            data={"csrf_token": csrf_token, "name": "Dyna Updated"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_user_admin(123456789).name == "Dyna Updated"

        response = client.post(
            "/admin/users/123456789/status",
            data={"csrf_token": csrf_token, "active": "0"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_user(123456789) is None
        assert database.get_user_admin(123456789).active is False

    actions = [row["action"] for row in database.list_admin_audit_events()]
    assert actions[:7] == [
        "user.deactivated",
        "user.updated",
        "membership.deactivated",
        "membership.updated",
        "membership.created",
        "membership.created",
        "user.created",
    ]


def test_company_instruction_draft_publish_and_runtime_override(tmp_path):
    companies_file = tmp_path / "companies.json"
    instruction_file = tmp_path / "companies" / "company-a" / "instruction.md"
    instruction_file.parent.mkdir(parents=True)
    instruction_file.write_text("Instruksi file lama", encoding="utf-8")
    companies_file.write_text(
        json.dumps(
            [
                {
                    "id": "company-a",
                    "name": "Company A",
                    "instruction_file": "companies/company-a/instruction.md",
                }
            ]
        ),
        encoding="utf-8",
    )
    database = Database(tmp_path / "instruction.db")
    database.initialize()
    database.bootstrap_companies(companies_file)
    company = database.get_company_admin("company-a")
    assert company is not None
    assert database.get_published_company_instruction("company-a") is None
    content = load_company_content(company, tmp_path, 1000)
    assert content.instruction == "Instruksi file lama"

    database.save_company_instruction_draft(
        "company-a", "Instruksi database v1", actor="admin"
    )
    assert database.get_published_company_instruction("company-a") is None
    assert database.publish_company_instruction("company-a", actor="admin") == 1
    published = database.get_published_company_instruction("company-a")
    assert published == "Instruksi database v1"
    content = load_company_content(company, tmp_path, 1000, published)
    assert content.instruction == "Instruksi database v1"

    database.save_company_instruction_draft(
        "company-a", "Instruksi database v2", actor="admin"
    )
    assert database.get_published_company_instruction("company-a") == (
        "Instruksi database v1"
    )
    assert database.publish_company_instruction("company-a", actor="admin") == 2
    versions = database.list_company_instruction_versions("company-a")
    assert [row["version_number"] for row in versions] == [2, 1]
    database.restore_company_instruction_version_to_draft(
        "company-a", versions[1]["id"], actor="admin"
    )
    state = database.get_company_instruction_admin("company-a")
    assert state is not None
    assert state["draft_content"] == "Instruksi database v1"
    assert state["published_content"] == "Instruksi database v2"


def test_admin_company_instruction_workflow(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Company A"}]),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text("[]", encoding="utf-8")
    database = Database(tmp_path / "instruction-admin.db")
    database.initialize()
    database.bootstrap_companies(companies_file)
    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "instruction-admin.db",
        project_root=Path("."),
    )
    app = create_admin_app(settings, database)
    with TestClient(app) as client:
        response = client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        response = client.get("/admin/instructions")
        assert response.status_code == 200
        assert "Company A" in response.text

        response = client.get("/admin/instructions/company-a")
        assert response.status_code == 200
        assert "Review untuk publish" in response.text
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', response.text)
        assert csrf is not None
        csrf_token = csrf.group(1)

        response = client.post(
            "/admin/instructions/company-a/draft",
            data={"content": "Jawab sesuai kebijakan Company A."},
        )
        assert response.status_code == 403
        response = client.post(
            "/admin/instructions/company-a/draft",
            data={
                "csrf_token": csrf_token,
                "content": "Jawab sesuai kebijakan Company A.",
                "submit_action": "preview",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert response.headers["location"] == (
            "/admin/instructions/company-a/preview"
        )
        assert database.get_published_company_instruction("company-a") is None

        response = client.get(response.headers["location"])
        assert response.status_code == 200
        assert "Jawab sesuai kebijakan Company A." in response.text
        response = client.post(
            "/admin/instructions/company-a/publish",
            data={"csrf_token": csrf_token},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_published_company_instruction("company-a") == (
            "Jawab sesuai kebijakan Company A."
        )
        response = client.get("/admin/instructions/company-a")
        assert "Published v1" in response.text
        assert "Live" in response.text

    actions = [row["action"] for row in database.list_admin_audit_events()]
    assert actions[:2] == [
        "company_instruction.published",
        "company_instruction.draft_saved",
    ]


def test_admin_module_playbook_and_membership_access_workflow(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Company A"}]),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text(
        json.dumps(
            [
                {
                    "telegram_id": 42,
                    "name": "DK",
                    "memberships": [
                        {
                            "company_id": "company-a",
                            "job_title": "Owner",
                            "division": "Management",
                            "role_level": "gm",
                            "communication_profile": "executive",
                            "default": True,
                        }
                    ],
                }
            ]
        ),
        encoding="utf-8",
    )
    database = Database(tmp_path / "module-admin.db")
    database.initialize()
    database.bootstrap_companies(companies_file)
    database.bootstrap_users(users_file)
    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "module-admin.db",
    )
    app = create_admin_app(settings, database)

    with TestClient(app) as client:
        assert client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        ).status_code == 303
        credential_page = client.get("/admin/runtime-profiles/new")
        assert credential_page.status_code == 200
        credential_csrf = re.search(
            r'name="csrf_token" value="([a-f0-9]+)"', credential_page.text
        )
        assert credential_csrf is not None
        csrf_token = credential_csrf.group(1)
        # Existing environment profiles remain usable after the new two-field UI.
        database.create_ai_runtime_profile(
            "marketing-openai", "Marketing OpenAI", "openai", "AI_KEY_MARKETING",
            "gpt-5.4-mini", "", actor="tester",
        )
        profile = database.get_ai_runtime_profile("marketing-openai")
        assert profile is not None
        assert profile.api_key_env == "AI_KEY_MARKETING"

        form_page = client.get("/admin/modules/new")
        assert form_page.status_code == 200
        assert 'name="module_id"' not in form_page.text
        csrf = re.search(
            r'name="csrf_token" value="([a-f0-9]+)"', form_page.text
        )
        assert csrf is not None
        csrf_token = csrf.group(1)

        assert client.post(
            "/admin/modules",
            data={
                "company_id": "company-a",
                "name": "Marketing",
                "description": "Campaign planning",
            },
        ).status_code == 403
        response = client.post(
            "/admin/modules",
            data={
                "csrf_token": csrf_token,
                "company_id": "company-a",
                "module_id": "forged-id",
                "name": "Marketing",
                "description": "Campaign planning",
                "ai_runtime_profile_id": "marketing-openai",
                "active": "1",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert response.headers["location"].startswith(
            "/admin/modules/company-a/marketing"
        )
        assert database.get_module_admin("company-a", "forged-id") is None
        saved_module = database.get_module_admin("company-a", "marketing")
        assert saved_module is not None
        assert saved_module.ai_runtime_profile_id == "marketing-openai"

        response = client.post(
            "/admin/modules/company-a/marketing/draft",
            data={
                "csrf_token": csrf_token,
                "content": "Playbook marketing yang hanya boleh muncul setelah publish.",
                "submit_action": "preview",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert response.headers["location"] == (
            "/admin/modules/company-a/marketing/preview"
        )
        assert database.get_published_module_playbook(
            "company-a", "marketing"
        ) is None
        preview = client.get(response.headers["location"])
        assert preview.status_code == 200
        assert "hanya boleh muncul setelah publish" in preview.text
        assert client.post(
            "/admin/modules/company-a/marketing/publish",
            data={"csrf_token": csrf_token},
            follow_redirects=False,
        ).status_code == 303

        membership_page = client.get(
            "/admin/users/42/memberships/company-a/edit"
        )
        assert membership_page.status_code == 200
        assert "Akses Modules" in membership_page.text
        assert 'value="marketing"' in membership_page.text
        assert client.post(
            "/admin/users/42/memberships/company-a/modules",
            data={"csrf_token": csrf_token, "module_ids": "marketing"},
            follow_redirects=False,
        ).status_code == 303
        assert [
            item.module_id
            for item in database.list_accessible_modules(42, "company-a")
        ] == ["marketing"]

        activity = client.get("/admin/activity?category=module")
        assert activity.status_code == 200
        assert "Playbook dipublikasikan" in activity.text
        assert "Akses module diperbarui" in activity.text
        assert "hanya boleh muncul setelah publish" not in activity.text


def test_knowledge_publish_runtime_fallback_and_tenant_isolation(tmp_path):
    companies_file = tmp_path / "companies.json"
    for company_id, file_content in (
        ("company-a", "Knowledge file A"),
        ("company-b", "Knowledge file B"),
    ):
        knowledge_dir = tmp_path / "companies" / company_id / "knowledge"
        knowledge_dir.mkdir(parents=True)
        (knowledge_dir / "legacy.md").write_text(file_content, encoding="utf-8")
    companies_file.write_text(
        json.dumps(
            [
                {
                    "id": "company-a",
                    "name": "Company A",
                    "knowledge_dir": "companies/company-a/knowledge",
                },
                {
                    "id": "company-b",
                    "name": "Company B",
                    "knowledge_dir": "companies/company-b/knowledge",
                },
            ]
        ),
        encoding="utf-8",
    )
    database = Database(tmp_path / "knowledge.db")
    database.initialize()
    database.bootstrap_companies(companies_file)
    company_a = database.get_company_admin("company-a")
    company_b = database.get_company_admin("company-b")
    assert company_a is not None and company_b is not None

    assert database.get_published_company_knowledge("company-a", 10_000) is None
    fallback_a = load_company_content(company_a, tmp_path, 10_000).knowledge
    assert "Knowledge file A" in fallback_a
    database.create_knowledge_document(
        "company-a", "sales-target", "Sales Target", "Target A v1", "admin"
    )
    assert database.get_published_company_knowledge("company-a", 10_000) is None
    assert database.publish_knowledge_document(
        "company-a", "sales-target", "admin"
    ) == 1
    published_a = database.get_published_company_knowledge("company-a", 10_000)
    assert published_a is not None
    assert "Target A v1" in published_a
    assert "Knowledge file A" not in published_a
    assert load_company_content(
        company_a, tmp_path, 10_000, None, published_a
    ).knowledge == published_a

    assert database.get_published_company_knowledge("company-b", 10_000) is None
    fallback_b = load_company_content(company_b, tmp_path, 10_000).knowledge
    assert "Knowledge file B" in fallback_b
    database.create_knowledge_document(
        "company-b", "sales-target", "Sales Target", "Target B", "admin"
    )
    database.publish_knowledge_document("company-b", "sales-target", "admin")
    assert "Target B" not in database.get_published_company_knowledge(
        "company-a", 10_000
    )

    database.save_knowledge_document_draft(
        "company-a", "sales-target", "Sales Target", "Target A v2", "admin"
    )
    assert "Target A v1" in database.get_published_company_knowledge(
        "company-a", 10_000
    )
    assert database.publish_knowledge_document(
        "company-a", "sales-target", "admin"
    ) == 2
    versions = database.list_knowledge_document_versions(
        "company-a", "sales-target"
    )
    assert [row["version_number"] for row in versions] == [2, 1]
    assert all(len(str(row["content_sha256"])) == 64 for row in versions)
    database.restore_knowledge_document_version_to_draft(
        "company-a", "sales-target", versions[1]["id"], "admin"
    )
    document = database.get_knowledge_document_admin(
        "company-a", "sales-target"
    )
    assert document is not None
    assert document["draft_content"] == "Target A v1"
    assert document["published_content"] == "Target A v2"

    try:
        database.restore_knowledge_document_version_to_draft(
            "company-b", "sales-target", versions[1]["id"], "admin"
        )
        assert False, "Cross-tenant restore seharusnya ditolak"
    except ValueError as exc:
        assert "tidak ditemukan" in str(exc)

    database.set_knowledge_document_active(
        "company-a", "sales-target", False, "admin"
    )
    assert database.get_published_company_knowledge("company-a", 10_000) == ""
    assert load_company_content(company_a, tmp_path, 10_000, None, "").knowledge == ""


def test_admin_knowledge_document_workflow(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Company A"}]),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text("[]", encoding="utf-8")
    database = Database(tmp_path / "knowledge-admin.db")
    database.initialize()
    database.bootstrap_companies(companies_file)
    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "knowledge-admin.db",
        project_root=Path("."),
    )
    app = create_admin_app(settings, database)
    with TestClient(app) as client:
        response = client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert client.get("/admin/knowledge").status_code == 200
        response = client.get("/admin/knowledge/company-a/new")
        assert response.status_code == 200
        assert 'name="document_key"' not in response.text
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', response.text)
        assert csrf is not None
        csrf_token = csrf.group(1)

        response = client.post(
            "/admin/knowledge/company-a",
            data={
                "document_key": "target-2026",
                "title": "Target 2026",
                "content": "Target omzet bulanan Rp500 juta.",
            },
        )
        assert response.status_code == 403
        response = client.post(
            "/admin/knowledge/company-a",
            data={
                "csrf_token": csrf_token,
                "document_key": "key-yang-dipalsukan",
                "title": "Target 2026",
                "content": "Target omzet bulanan Rp500 juta.",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert response.headers["location"].startswith(
            "/admin/knowledge/company-a/target-2026"
        )
        assert database.get_published_company_knowledge("company-a", 10_000) is None

        editor = client.get("/admin/knowledge/company-a/target-2026")
        assert "Review untuk publish" in editor.text
        response = client.post(
            "/admin/knowledge/company-a/target-2026/draft",
            data={
                "csrf_token": csrf_token,
                "title": "Target 2026 Final",
                "content": "Target omzet bulanan Rp550 juta.",
                "submit_action": "preview",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert response.headers["location"] == (
            "/admin/knowledge/company-a/target-2026/preview"
        )
        response = client.get(response.headers["location"])
        assert response.status_code == 200
        assert "Target 2026 Final" in response.text
        assert "Rp550 juta" in response.text
        response = client.post(
            "/admin/knowledge/company-a/target-2026/publish",
            data={"csrf_token": csrf_token},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert "Rp550 juta" in database.get_published_company_knowledge(
            "company-a", 10_000
        )
        response = client.get("/admin/knowledge/company-a/target-2026")
        assert response.status_code == 200
        assert "Published v1" in response.text
        assert "Live" in response.text

        response = client.post(
            "/admin/knowledge/company-a/target-2026/status",
            data={"csrf_token": csrf_token, "active": "0"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert database.get_published_company_knowledge("company-a", 10_000) == ""

    actions = [row["action"] for row in database.list_admin_audit_events()]
    assert actions[:4] == [
        "knowledge_document.deactivated",
        "knowledge_document.published",
        "knowledge_document.draft_saved",
        "knowledge_document.created",
    ]
    audit_payload = " ".join(
        str(row["details_json"]) for row in database.list_admin_audit_events()
    )
    assert "Target omzet bulanan Rp550 juta." not in audit_payload


def test_document_ingestion_pdf_docx_and_invalid_formats():
    pdf_buffer = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(
        b"BT /F1 12 Tf 72 720 Td (Target PDF Rp700 juta) Tj ET"
    )
    page[NameObject("/Contents")] = content
    writer.write(pdf_buffer)
    pdf = extract_uploaded_document(
        "target.pdf", "application/pdf", pdf_buffer.getvalue()
    )
    assert "Target PDF Rp700 juta" in pdf.text
    assert len(pdf.sha256) == 64

    docx_buffer = BytesIO()
    document = Document()
    document.add_heading("Brand Identity", level=1)
    document.add_paragraph("Amazing Malang memakai tone hangat dan informatif.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Warna"
    table.cell(0, 1).text = "Ungu"
    document.save(docx_buffer)
    docx = extract_uploaded_document(
        "brand.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        docx_buffer.getvalue(),
    )
    assert "Brand Identity" in docx.text
    assert "Warna | Ungu" in docx.text

    for filename, data, expected in (
        ("legacy.doc", b"legacy", "Simpan ulang sebagai .docx"),
        ("fake.pdf", b"not-a-pdf", "tidak cocok dengan format PDF"),
        ("scan.pdf", _blank_pdf_bytes(), "tidak mengandung teks"),
    ):
        try:
            extract_uploaded_document(filename, "application/octet-stream", data)
            assert False, f"{filename} seharusnya ditolak"
        except ValueError as exc:
            assert expected in str(exc)


def test_admin_knowledge_docx_upload_creates_reviewable_draft(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Company A"}]),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text("[]", encoding="utf-8")
    database = Database(tmp_path / "upload-admin.db")
    database.initialize()
    database.bootstrap_companies(companies_file)
    settings = _test_settings(
        tmp_path,
        users_file,
        companies_file,
        database_path=tmp_path / "upload-admin.db",
        project_root=Path("."),
    )
    app = create_admin_app(settings, database)
    docx_buffer = BytesIO()
    source_document = Document()
    source_document.add_paragraph("Brand voice Amazing Malang bersifat hangat.")
    source_document.save(docx_buffer)
    payload = docx_buffer.getvalue()

    with TestClient(app) as client:
        assert client.post(
            "/admin/login",
            data={"username": "admin", "password": "strong-password"},
            follow_redirects=False,
        ).status_code == 303
        form_page = client.get("/admin/knowledge/company-a/new")
        csrf = re.search(
            r'name="csrf_token" value="([a-f0-9]+)"', form_page.text
        )
        assert csrf is not None
        response = client.post(
            "/admin/knowledge/company-a",
            data={
                "csrf_token": csrf.group(1),
                "document_key": "key-yang-dipalsukan",
                "title": "Brand Identity",
                "content": "",
            },
            files={
                "source_file": (
                    "brand.docx",
                    payload,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        stored = database.get_knowledge_document_admin(
            "company-a", "brand-identity"
        )
        assert stored is not None
        assert "Brand voice Amazing Malang" in stored["draft_content"]
        assert stored["source_filename"] == "brand.docx"
        assert stored["source_size_bytes"] == len(payload)
        assert len(str(stored["source_sha256"])) == 64
        assert database.get_published_company_knowledge("company-a", 10_000) is None
        editor = client.get("/admin/knowledge/company-a/brand-identity")
        assert "Source upload" in editor.text
        assert "brand.docx" in editor.text

    audit = database.list_admin_audit_events()[0]
    assert "brand.docx" in str(audit["details_json"])
    assert "Brand voice Amazing Malang" not in str(audit["details_json"])


def _blank_pdf_bytes() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(output)
    return output.getvalue()


def test_bot_serializes_same_user_and_allows_different_users(tmp_path):
    companies_file = tmp_path / "companies.json"
    companies_file.write_text(
        json.dumps([{"id": "company-a", "name": "Company A"}]),
        encoding="utf-8",
    )
    users_file = tmp_path / "users.json"
    users_file.write_text(
        json.dumps(
            [
                {
                    "telegram_id": 42,
                    "name": "User A",
                    "memberships": [{"company_id": "company-a", "default": True}],
                },
                {
                    "telegram_id": 43,
                    "name": "User B",
                    "memberships": [{"company_id": "company-a", "default": True}],
                },
            ]
        ),
        encoding="utf-8",
    )
    database = Database(tmp_path / "concurrency.db")
    database.initialize()
    database.sync_companies(companies_file)
    database.sync_users(users_file)
    settings = _test_settings(tmp_path, users_file, companies_file)
    bot = InternalBot(settings, database, SimpleNamespace())

    async def measure(user_ids: list[int]) -> int:
        active = 0
        max_active = 0

        async def worker(user_id: int) -> None:
            nonlocal active, max_active
            update = SimpleNamespace(
                effective_user=SimpleNamespace(id=user_id),
                effective_message=None,
            )
            async with bot._serialized_user(update) as user:
                assert user is not None
                active += 1
                max_active = max(max_active, active)
                await asyncio.sleep(0.02)
                active -= 1

        await asyncio.gather(*(worker(user_id) for user_id in user_ids))
        return max_active

    async def scenario() -> tuple[int, int]:
        return await measure([42, 42]), await measure([42, 43])

    same_user_max, different_user_max = asyncio.run(scenario())
    assert same_user_max == 1
    assert different_user_max == 2

    application = bot.build_application()
    assert application.update_processor.max_concurrent_updates == 4


def test_split_message():
    chunks = _split_message("a" * 9000, size=4000)
    assert [len(chunk) for chunk in chunks] == [4000, 4000, 1000]


def test_split_plain_text_preserves_paragraphs_and_indentation():
    source = to_telegram_plain_text("```python\n" + "# comment\n" * 360 + "    print('ok')\n```")
    chunks = _split_message(source, size=3500)
    assert "".join(chunks) == source
    assert all(len(chunk) <= 3500 for chunk in chunks)
    assert "    print('ok')" in chunks[-1]


def test_telegram_renderer_removes_legacy_markup():
    source = (
        "# Diagnosis\n\n"
        "**Conversion** turun dan *perlu dicek*.\n"
        "- Buka `/whoami`\n"
        "- Baca [panduan](https://example.com/guide?a=1&b=2)\n"
        "> Data belum cukup\n"
        "||Jawaban quiz||"
    )
    assert to_telegram_plain_text(source) == (
        "Diagnosis\n\nConversion turun dan perlu dicek.\n"
        "- Buka /whoami\n- Baca panduan (https://example.com/guide?a=1&b=2)\n"
        "Data belum cukup\nJawaban quiz"
    )


def test_telegram_plain_text_preserves_literal_html_and_code():
    source = "<b>raw</b> & aman\n```python\nif a < b:\n    print('&')\n```"
    assert to_telegram_plain_text(source) == "<b>raw</b> & aman\nif a < b:\n    print('&')"


def test_telegram_plain_text_threads_copy_paste_regression():
    source = (
        "1. Edukasi Farm\n> Batu dingin gini enaknya ngapain\n>\n"
        "> Aku sama kawan lagi demen jemur pagi 🐐\n\n"
        "**Pilih 1 yang paling pas buat diposting hari ini**\n"
        "Ketik *belum* atau `/module threads-generator`."
    )
    assert to_telegram_plain_text(source) == (
        "1. Edukasi Farm\nBatu dingin gini enaknya ngapain\n\n"
        "Aku sama kawan lagi demen jemur pagi 🐐\n\n"
        "Pilih 1 yang paling pas buat diposting hari ini\n"
        "Ketik belum atau /module threads-generator."
    )


@pytest.mark.parametrize(("source", "expected"), [
    ("#MalangStrudel @piko harga Rp25.000 & diskon 5% 🐐", "#MalangStrudel @piko harga Rp25.000 & diskon 5% 🐐"),
    ("api_key user_name x > 3\n>= 10\n>100\n2 * 3 * 4", "api_key user_name x > 3\n>= 10\n>100\n2 * 3 * 4"),
    ("***tebal miring*** __tebal__ _miring_ ~~hapus~~ ||spoiler||", "tebal miring tebal miring hapus spoiler"),
    ("**tebal dan *miring***", "tebal dan miring"),
    ("[**Panduan**](https://example.com/a_(b)?x=1&y=2)", "Panduan (https://example.com/a_(b)?x=1&y=2)"),
    ("[https://example.com](https://example.com)", "https://example.com"),
    ("https://example.com/__a__?x=1&y=2", "https://example.com/__a__?x=1&y=2"),
    (r"\*literal\* \`literal\`", "*literal* `literal`"),
    (r"`\*literal\*`", r"\*literal\*"),
    ("```python\n# comment\nx = '**literal**'\n    y = 2 ** 3\n```", "# comment\nx = '**literal**'\n    y = 2 ** 3"),
    ("> > nested\n>\n> paragraf\r\n* pilihan", "nested\n\nparagraf\n- pilihan"),
    ("\x00plain0\x00 `aman`", "\x00plain0\x00 aman"),
    (">\n---\n** **", ""),
])
def test_telegram_plain_text_keeps_content(source, expected):
    assert to_telegram_plain_text(source) == expected


@pytest.mark.parametrize(("source", "expected"), [
    ("**Analisis Threads**\n> Kutipan biasa", [TelegramResponsePart("**Analisis Threads**\n> Kutipan biasa")]),
    ("**Pilih akun**\n- a. Piko\n- b. Farm", [TelegramResponsePart("**Pilih akun**\n- a. Piko\n- b. Farm")]),
    ("**Opsi 1**\n[[COPY_TEXT]]\n> Caption **Piko**\n>\n> #Malang 🐐\n[[/COPY_TEXT]]\n**Pilih yang sesuai**", [
        TelegramResponsePart("**Opsi 1**"), TelegramResponsePart("Caption Piko\n\n#Malang 🐐", True), TelegramResponsePart("**Pilih yang sesuai**"),
    ]),
    ("[[COPY_TEXT]]\nCaption A\n[[/COPY_TEXT]]\n[[COPY_TEXT]]\nCaption B\n[[/COPY_TEXT]]", [TelegramResponsePart("Caption A", True), TelegramResponsePart("Caption B", True)]),
    ("[[COPY_TEXT]]\n**Caption belum ditutup**", [TelegramResponsePart("Caption belum ditutup", True)]),
    ("[[/COPY_TEXT]]\n**Analisis**", [TelegramResponsePart("**Analisis**")]),
    ("[[COPY_TEXT]]\n[[COPY_TEXT]]\nIsi\n[[/COPY_TEXT]]", [TelegramResponsePart("Isi", True)]),
    ("[[COPY_TEXT]]\n>\n---\n[[/COPY_TEXT]]", []),
    ("Contoh penanda [[COPY_TEXT]] sebagai teks.", [TelegramResponsePart("Contoh penanda [[COPY_TEXT]] sebagai teks.")]),
    ("```text\n[[COPY_TEXT]]\nContoh literal\n[[/COPY_TEXT]]\n```", [TelegramResponsePart("```text\n[[COPY_TEXT]]\nContoh literal\n[[/COPY_TEXT]]\n```")]),
    ("[[COPY_TEXT]]\r\nParagraf 1\r\n\r\nParagraf 2\r\n[[/COPY_TEXT]]", [TelegramResponsePart("Paragraf 1\n\nParagraf 2", True)]),
])
def test_copy_ready_parts_do_not_flatten_other_answers(source, expected):
    assert prepare_response_parts(source, 12000) == expected


@pytest.mark.parametrize("limit", [0, 1, 7, 13, 40, 12000])
def test_copy_ready_parts_share_budget_without_leaking_markers(limit):
    source = "**Judul**\n[[COPY_TEXT]]\n" + "caption " * 1000 + "\n[[/COPY_TEXT]]\n**Tips**"
    parts = prepare_response_parts(source, limit)
    history = "\n\n".join(part.text for part in parts)
    assert len(history) <= limit
    assert "[[COPY_TEXT]]" not in history and "[[/COPY_TEXT]]" not in history
    if len(parts) > 1:
        assert parts[1].copyable is True


def test_ordinary_html_renderer_keeps_format_and_escapes_raw_html():
    source = "# Diagnosis\n**tebal** *miring* `command`\n> quote\n||spoiler||\n<b>literal</b> & data"
    rendered = markdown_to_telegram_html(source)
    for fragment in ("<b>Diagnosis</b>", "<b>tebal</b>", "<i>miring</i>", "<code>command</code>", "<blockquote>quote</blockquote>", "<tg-spoiler>spoiler</tg-spoiler>", "&lt;b&gt;literal&lt;/b&gt; &amp; data"):
        assert fragment in rendered


def test_empty_plain_text_answer_does_not_write_history(tmp_path):
    database = Database(tmp_path / "empty-response.db")
    database.initialize()
    database.create_company("company-a", "Company A", "admin")
    database.create_user_with_membership(42, "DK", "company-a", "Owner", "", "gm", "executive", "", "admin")
    replies = []

    async def generate(*_args):
        return "[[COPY_TEXT]]\n>\n---\n```\n\n```\n[[/COPY_TEXT]]"

    async def reply_text(text, **_kwargs):
        replies.append(text)

    async def send_chat_action(*_args, **_kwargs):
        pass

    settings = _test_settings(tmp_path, tmp_path / "users.json", tmp_path / "companies.json", project_root=tmp_path)
    bot = InternalBot(settings, database, SimpleNamespace(generate=generate))
    message = SimpleNamespace(text="belum", reply_text=reply_text)
    update = SimpleNamespace(effective_user=SimpleNamespace(id=42), effective_chat=SimpleNamespace(id=42), effective_message=message)
    context = SimpleNamespace(bot=SimpleNamespace(send_chat_action=send_chat_action))
    asyncio.run(bot.chat(update, context))
    assert len(replies) == 1 and "gagal memproses" in replies[0]
    assert database.get_history(42, 12, "company-a", "") == []


def _ai_pair():
    primary = AIRuntimeProfile("primary", "AI Primary", "openai", "TEST_AI_PRIMARY", "test-primary", "", True)
    backup = AIRuntimeProfile("backup", "AI Backup", "deepseek", "TEST_AI_BACKUP", "test-backup", "", True)
    return primary, backup


def _api_failure(status):
    return APIStatusError(
        "private-provider-body", response=httpx.Response(
            status, request=httpx.Request("POST", "https://example.invalid/secret-url")
        ), body={"secret": "do-not-log"},
    )


@pytest.mark.parametrize("failure", ["connection", "timeout", 408, 429, 500, 502, 503, 504])
def test_module_failover_sends_same_context_once_and_redacts_errors(monkeypatch, caplog, failure):
    primary, backup = _ai_pair()
    monkeypatch.setenv(primary.api_key_env, "private-primary-key")
    monkeypatch.setenv(backup.api_key_env, "private-backup-key")
    calls = []
    backup_loads = []

    class Provider:
        def __init__(self, model):
            self.model = model

        async def generate(self, prompt, history, user_text):
            calls.append((self.model, prompt, list(history), user_text))
            if self.model == primary.model:
                if failure == "connection":
                    raise APIConnectionError(request=httpx.Request("POST", "https://example.invalid"))
                if failure == "timeout":
                    raise TimeoutError("private-provider-body")
                raise _api_failure(failure)
            return "Jawaban cadangan"

    resolver = ModuleProviderResolver(lambda _p, _k, model, _u: Provider(model))
    history = [{"role": "user", "content": "history-company-a"}]

    def load_backup():
        backup_loads.append(True)
        return backup

    answer = asyncio.run(resolver.generate(primary, load_backup, "private-prompt-a", history, "private-question"))
    assert answer == "Jawaban cadangan"
    assert len(calls) == 2
    assert calls[0][1:] == calls[1][1:] == ("private-prompt-a", history, "private-question")
    assert [call[0] for call in calls] == [primary.model, backup.model]
    assert backup_loads == [True]
    assert "module_ai_failover" in caplog.text
    for value in ("private-provider-body", "do-not-log", "private-primary-key", "private-backup-key", "private-prompt-a", "private-question", "secret-url"):
        assert value not in caplog.text


@pytest.mark.parametrize("failure", [400, 401, 403, 404, 409, 422, "unexpected"])
def test_module_does_not_failover_on_permanent_errors(monkeypatch, failure):
    primary, _ = _ai_pair()
    monkeypatch.setenv(primary.api_key_env, "private-key")

    async def generate(*_args):
        if failure == "unexpected":
            raise RuntimeError("private-error")
        raise _api_failure(failure)

    resolver = ModuleProviderResolver(lambda *_args: SimpleNamespace(generate=generate))
    with pytest.raises(ModuleGenerationError) as raised:
        asyncio.run(resolver.generate(primary, lambda: pytest.fail("Backup must not be loaded"), "", [], ""))
    assert "private" not in str(raised.value)


def test_module_success_never_resolves_backup_and_next_request_tries_primary(monkeypatch):
    primary, backup = _ai_pair()
    monkeypatch.setenv(primary.api_key_env, "primary-key")
    monkeypatch.setenv(backup.api_key_env, "backup-key")
    calls = []

    def factory(_provider, _key, model, _url):
        async def generate(*_args):
            calls.append(model)
            if len(calls) == 1:
                raise _api_failure(503)
            return model
        return SimpleNamespace(generate=generate)

    resolver = ModuleProviderResolver(factory)
    assert asyncio.run(resolver.generate(primary, lambda: backup, "", [], "")) == backup.model
    assert asyncio.run(resolver.generate(primary, lambda: pytest.fail("No fallback on success"), "", [], "")) == primary.model
    assert calls == [primary.model, backup.model, primary.model]


@pytest.mark.parametrize("invalid_primary", ["missing-key", "inactive", "missing-profile"])
def test_module_bad_primary_configuration_does_not_use_backup(monkeypatch, invalid_primary):
    primary, backup = _ai_pair()
    monkeypatch.setenv(backup.api_key_env, "backup-key")
    monkeypatch.delenv(primary.api_key_env, raising=False)
    if invalid_primary == "inactive":
        primary = replace(primary, active=False)
    elif invalid_primary == "missing-profile":
        primary = None
    resolver = ModuleProviderResolver(lambda *_args: pytest.fail("Must not call any provider"))
    with pytest.raises(RuntimeCredentialError):
        asyncio.run(resolver.generate(primary, lambda: pytest.fail("Must not load backup"), "", [], ""))


@pytest.mark.parametrize("backup_state", ["none", "inactive", "missing-key", "same", "fails"])
def test_module_backup_failure_stops_after_at_most_two_attempts(monkeypatch, backup_state):
    primary, backup = _ai_pair()
    monkeypatch.setenv(primary.api_key_env, "primary-key")
    monkeypatch.setenv(backup.api_key_env, "backup-key")
    calls = []

    async def generate(*_args):
        calls.append(True)
        raise _api_failure(503)

    if backup_state == "none":
        backup = None
    elif backup_state == "inactive":
        backup = replace(backup, active=False)
    elif backup_state == "missing-key":
        monkeypatch.delenv(backup.api_key_env)
    elif backup_state == "same":
        backup = primary
    resolver = ModuleProviderResolver(lambda *_args: SimpleNamespace(generate=generate))
    with pytest.raises((RuntimeCredentialError, ModuleGenerationError)):
        asyncio.run(resolver.generate(primary, lambda: backup, "", [], ""))
    assert len(calls) == (2 if backup_state == "fails" else 1)


def test_module_timeout_and_cancellation_are_bounded(monkeypatch):
    primary, backup = _ai_pair()
    monkeypatch.setenv(primary.api_key_env, "primary-key")
    monkeypatch.setenv(backup.api_key_env, "backup-key")
    monkeypatch.setattr("app.providers.MODULE_AI_TIMEOUT_SECONDS", 0.01)
    calls = []

    def factory(_provider, _key, model, _url):
        async def generate(*_args):
            calls.append(model)
            if model == primary.model:
                await asyncio.Event().wait()
            return "backup-answer"
        return SimpleNamespace(generate=generate)

    resolver = ModuleProviderResolver(factory)
    assert asyncio.run(resolver.generate(primary, lambda: backup, "", [], "")) == "backup-answer"
    assert calls == [primary.model, backup.model]

    async def cancelled(*_args):
        raise asyncio.CancelledError()

    resolver = ModuleProviderResolver(lambda *_args: SimpleNamespace(generate=cancelled))
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(resolver.generate(primary, lambda: pytest.fail("Cancellation is not fallback"), "", [], ""))


def _database_with_ai_pair(tmp_path):
    database = Database(tmp_path / "ai-pair.db")
    database.initialize()
    database.create_company("company-a", "Company A", "admin")
    primary, backup = _ai_pair()
    for profile in (primary, backup):
        database.create_ai_runtime_profile(
            profile.profile_id, profile.label, profile.provider,
            profile.api_key_env, profile.model, profile.base_url, "admin",
        )
    return database, primary, backup


def test_module_ai_pair_persistence_validation_and_legacy_migration(tmp_path):
    database, primary, backup = _database_with_ai_pair(tmp_path)
    old_module = database.create_module("company-a", "legacy", "Legacy", "", "admin", primary.profile_id)
    assert old_module.backup_ai_runtime_profile_id == ""
    # Simulate the last deployed schema, without the new nullable column.
    with database._connect() as connection:
        connection.execute("ALTER TABLE modules DROP COLUMN backup_ai_runtime_profile_id")
    database.initialize()
    database.initialize()
    assert database.get_module_admin("company-a", "legacy") == old_module
    database.update_module("company-a", "legacy", "Legacy", "", "admin", primary.profile_id, backup.profile_id)
    saved = database.get_module_admin("company-a", "legacy")
    assert saved.backup_ai_runtime_profile_id == backup.profile_id
    database.update_module("company-a", "legacy", "Legacy", "", "admin", primary.profile_id)
    assert database.get_module_admin("company-a", "legacy") == saved
    assert database.list_modules_admin()[0]["backup_ai_label"] == backup.label
    assert database.get_module_playbook_admin("company-a", "legacy")["backup_ai_label"] == backup.label
    for invalid in (primary.profile_id, "does-not-exist", "../invalid"):
        with pytest.raises(ValueError):
            database.update_module("company-a", "legacy", "Changed", "", "admin", primary.profile_id, invalid)
        assert database.get_module_admin("company-a", "legacy") == saved
        with pytest.raises(ValueError):
            database.create_module("company-a", "new", "New", "", "admin", primary.profile_id, backup_ai_runtime_profile_id=invalid)
        assert database.get_module_admin("company-a", "new") is None
    database.set_ai_runtime_profile_active(backup.profile_id, False, "admin")
    with pytest.raises(ValueError):
        database.update_module("company-a", "legacy", "Changed", "", "admin", primary.profile_id, backup.profile_id)
    # A disabled backup does not disable the primary; it can be explicitly removed.
    database.update_module("company-a", "legacy", "Legacy", "", "admin", primary.profile_id, "")
    assert database.get_module_admin("company-a", "legacy").backup_ai_runtime_profile_id == ""


def test_admin_module_ai_pair_create_update_and_csrf(tmp_path):
    database, primary, backup = _database_with_ai_pair(tmp_path)
    settings = _test_settings(tmp_path, tmp_path / "users.json", tmp_path / "companies.json", database_path=database.path)
    with TestClient(create_admin_app(settings, database)) as client:
        assert client.post("/admin/modules", data={}, follow_redirects=False).status_code == 303
        client.post("/admin/login", data={"username": "admin", "password": "strong-password"})
        page = client.get("/admin/modules/new")
        assert "AI utama" in page.text and "AI cadangan" in page.text
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', page.text).group(1)
        data = {"csrf_token": csrf, "company_id": "company-a", "name": "Marketing", "active": "1", "ai_runtime_profile_id": primary.profile_id, "backup_ai_runtime_profile_id": backup.profile_id}
        assert client.post("/admin/modules", data={**data, "csrf_token": "bad"}).status_code == 403
        assert client.post("/admin/modules", data={**data, "backup_ai_runtime_profile_id": primary.profile_id}).status_code == 400
        assert client.post("/admin/modules", data=data, follow_redirects=False).status_code == 303
        saved = database.get_module_admin("company-a", "marketing")
        assert saved.backup_ai_runtime_profile_id == backup.profile_id
        path = "/admin/modules/company-a/marketing"
        editor = client.get(path)
        assert f'value="{backup.profile_id}" selected' in editor.text
        assert "AI terdaftar" in client.get("/admin/modules").text
        assert client.post(path, data={**data, "csrf_token": "bad"}).status_code == 403
        client.post(path, data={**data, "ai_runtime_profile_id": backup.profile_id, "backup_ai_runtime_profile_id": primary.profile_id})
        changed = database.get_module_admin("company-a", "marketing")
        assert changed.ai_runtime_profile_id == backup.profile_id
        assert changed.backup_ai_runtime_profile_id == primary.profile_id
        client.post(path, data={**data, "backup_ai_runtime_profile_id": "missing"})
        assert database.get_module_admin("company-a", "marketing") == changed
        database.set_ai_runtime_profile_active(primary.profile_id, False, "admin")
        assert f'value="{primary.profile_id}" selected disabled' in client.get(path).text
        client.post(path, data={**data, "ai_runtime_profile_id": backup.profile_id, "backup_ai_runtime_profile_id": ""})
        assert database.get_module_admin("company-a", "marketing").backup_ai_runtime_profile_id == ""


def test_module_sdk_has_single_attempt_without_changing_general(monkeypatch):
    from app.providers import create_provider
    primary, _ = _ai_pair()
    monkeypatch.setenv(primary.api_key_env, "fake-key")
    module_provider = ModuleProviderResolver().resolve(primary)
    global_provider = create_provider("openai", "fake-key", "test", None)
    assert module_provider.client.max_retries == 0
    assert module_provider.client.timeout == 30.0
    assert global_provider.client.max_retries == 2

    async def close():
        await module_provider.client.close()
        await global_provider.client.close()
    asyncio.run(close())


@pytest.mark.parametrize("mode", ["primary", "backup", "both-fail", "no-access"])
def test_bot_failover_keeps_tenant_history_and_authorization(tmp_path, monkeypatch, mode):
    database, primary, backup = _database_with_ai_pair(tmp_path)
    monkeypatch.setenv(primary.api_key_env, "private-primary")
    monkeypatch.setenv(backup.api_key_env, "private-backup")
    database.create_company("company-b", "Company B", "admin")
    database.create_user_with_membership(42, "DK", "company-a", "Owner", "Management", "gm", "executive", "", "admin")
    database.create_module("company-a", "marketing", "Marketing", "", "admin", primary.profile_id, backup_ai_runtime_profile_id=backup.profile_id)
    database.save_module_playbook_draft("company-a", "marketing", "playbook-company-a", "admin")
    database.publish_module_playbook("company-a", "marketing", "admin")
    database.set_membership_module_access(42, "company-a", ["marketing"], "admin")
    database.set_active_module(42, "marketing")
    database.create_knowledge_document("company-a", "facts", "Facts A", "knowledge-company-a", "admin")
    database.publish_knowledge_document("company-a", "facts", "admin")
    database.create_knowledge_document("company-b", "facts", "Facts B", "private-knowledge-company-b", "admin")
    database.publish_knowledge_document("company-b", "facts", "admin")
    database.add_message(42, "user", "general-only", "company-a", "")
    database.add_message(42, "user", "private-history-company-b", "company-b", "marketing")
    database.add_message(42, "user", "module-history", "company-a", "marketing")
    before = database.get_history(42, 20, "company-a", "marketing")
    calls = []
    replies = []
    global_calls = []

    def factory(_provider, _key, model, _url):
        async def generate(prompt, history, user_text):
            calls.append((model, prompt, list(history), user_text))
            if mode == "both-fail" or (mode == "backup" and model == primary.model):
                raise _api_failure(503)
            return "module-answer"
        return SimpleNamespace(generate=generate)

    async def global_generate(*_args):
        global_calls.append(True)
        return "general-answer"

    async def reply_text(text, **_kwargs):
        replies.append(text)

    async def send_chat_action(*_args, **_kwargs):
        pass

    settings = _test_settings(tmp_path, tmp_path / "users.json", tmp_path / "companies.json", database_path=database.path, project_root=tmp_path)
    bot = InternalBot(settings, database, SimpleNamespace(generate=global_generate), ModuleProviderResolver(factory))
    update = SimpleNamespace(effective_user=SimpleNamespace(id=42), effective_chat=SimpleNamespace(id=42), effective_message=SimpleNamespace(text="new-question", reply_text=reply_text))
    context = SimpleNamespace(bot=SimpleNamespace(send_chat_action=send_chat_action))
    if mode == "no-access":
        database.set_membership_module_access(42, "company-a", [], "admin")
    asyncio.run(bot.chat(update, context))
    after = database.get_history(42, 20, "company-a", "marketing")
    if mode == "no-access":
        assert calls == [] and global_calls == [True]
        assert after == before
        return
    assert global_calls == []
    assert len(calls) == (1 if mode == "primary" else 2)
    for _model, prompt, history, text in calls:
        assert "playbook-company-a" in prompt and "knowledge-company-a" in prompt
        assert "private-knowledge-company-b" not in prompt
        assert history == before
        assert text == "new-question"
    if mode == "both-fail":
        assert after == before
        assert "tidak tersedia" in replies[-1]
    else:
        assert after == before + [{"role": "user", "content": "new-question"}, {"role": "assistant", "content": "module-answer"}]
        assert replies == ["module-answer"]


@pytest.mark.parametrize("response_kind", ["retryable", "empty", "refusal"])
def test_module_failover_through_real_sdk_with_mock_http(monkeypatch, response_kind):
    from openai import AsyncOpenAI

    primary, backup = _ai_pair()
    monkeypatch.setenv(primary.api_key_env, "fake-primary-key")
    monkeypatch.setenv(backup.api_key_env, "fake-backup-key")
    requests = []
    clients = []

    def handle(request):
        payload = json.loads(request.content)
        requests.append(payload)
        if payload["model"] == primary.model and response_kind == "retryable":
            return httpx.Response(503, json={"error": {"message": "unavailable"}})
        message = {"role": "assistant", "content": "Mock HTTP answer"}
        if response_kind == "empty":
            message["content"] = "  "
        elif response_kind == "refusal":
            message.update(content=None, refusal="Cannot fulfill this request")
        return httpx.Response(200, json={"choices": [{"index": 0, "message": message, "finish_reason": "stop"}]})

    def client_factory(**kwargs):
        client = AsyncOpenAI(**kwargs)
        clients.append(client)
        return client

    monkeypatch.setattr("app.providers.AsyncClient", lambda **kwargs: httpx.AsyncClient(**kwargs, transport=httpx.MockTransport(handle)))
    monkeypatch.setattr("app.providers.AsyncOpenAI", client_factory)

    async def scenario():
        resolver = ModuleProviderResolver()
        try:
            if response_kind == "retryable":
                assert await resolver.generate(primary, lambda: backup, "system", [], "question") == "Mock HTTP answer"
            else:
                with pytest.raises(ModuleGenerationError):
                    await resolver.generate(primary, lambda: pytest.fail("Empty/refusal must not trigger backup"), "system", [], "question")
        finally:
            for client in clients:
                await client.close()

    asyncio.run(scenario())
    assert len(requests) == (2 if response_kind == "retryable" else 1)
    if response_kind == "retryable":
        assert [request["model"] for request in requests] == [primary.model, backup.model]
        assert requests[0]["messages"] == requests[1]["messages"]
