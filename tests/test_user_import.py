import asyncio
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
import re
import sqlite3
import struct
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from starlette.requests import Request

from app.admin import _bounded_import_request, create_admin_app
from app.database import Database
from app.user_import import (
    HEADERS, MAX_USER_IMPORT_BYTES, UserImportRow,
    UserImportValidationError, read_user_import,
)
from test_core import _test_settings


def xlsx(rows, *, sheet_name="Data_User"):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def xls(rows):
    """Small real BIFF8 workbook fixture, using text LABEL cells (no writer dependency)."""
    def record(opcode, payload=b""):
        return struct.pack("<HH", opcode, len(payload)) + payload

    def bof(kind):
        return record(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 1997, 0, 6))

    name = b"Data_User"
    globals_start = bof(5) + record(0x0042, struct.pack("<H", 1200))
    bound_size = 4 + 8 + len(name)
    offset = len(globals_start) + bound_size + 4
    bound = record(0x0085, struct.pack("<IBBBB", offset, 0, 0, len(name), 0) + name)
    sheet = bof(0x0010)
    sheet += record(0x0200, struct.pack("<IIHHH", 0, len(rows), 0, 5, 0))
    for row_number, values in enumerate(rows):
        for column, value in enumerate(values):
            text = str(value)
            sheet += record(0x0204, struct.pack("<HHHHB", row_number, column, 0, len(text), 1) + text.encode("utf-16-le"))
    return globals_start + bound + record(0x000A) + sheet + record(0x000A)


@pytest.fixture
def database(tmp_path):
    db = Database(tmp_path / "import.db")
    db.initialize()
    db.create_company("company-a", "Company A", "admin")
    db.create_company("company-b", "Company B", "admin")
    return db


def row(number=2, telegram_id=123456789, company="Company A", name="New User", **kwargs):
    return UserImportRow(number, name, telegram_id, company,
                         kwargs.get("job_title", "Marketing Manager"), kwargs.get("division", "Marketing"))


def run_import(database, rows, **kwargs):
    return database.import_new_users(rows, role_level="staff", communication_profile="staff",
                                     active=kwargs.get("active", False), actor="admin")


@pytest.mark.parametrize("writer, extension", [(xlsx, ".xlsx"), (xls, ".xls")])
def test_both_excel_formats(writer, extension):
    result = read_user_import("users" + extension, writer([HEADERS, ["Ajeng", "106545875", "Company A", "Manager", "Marketing"]]))
    assert result == [UserImportRow(2, "Ajeng", "106545875", "Company A", "Manager", "Marketing")]


def test_parser_blank_rows_numeric_id_reordered_headers_and_selected_tab():
    payload = xlsx([
        ["telegram id", "nama", "perusahaan", "divisi", "jabatan"],
        [None] * 5,
        [123456789, "Ajeng", "Company A", "Marketing", "Manager"],
    ])
    result = read_user_import("users.XLSX", payload)
    assert result[0].row_number == 3 and result[0].telegram_id == 123456789
    assert result[0].division == "Marketing"
    book = openpyxl.load_workbook(BytesIO(payload))
    book.create_sheet("Contoh").append(["not a header"])
    book.active = 1
    output = BytesIO()
    book.save(output)
    assert read_user_import("multi.xlsx", output.getvalue()) == result


@pytest.mark.parametrize("payload, message", [
    ([], "header"),
    ([HEADERS], "Belum ada"),
    ([["Nama", "Telegram ID", "Perusahaan", "Jabatan"]], "lima kolom"),
    ([["Nama", "Telegram ID", "Perusahaan", "Jabatan", "Nama"]], "lima kolom"),
    ([HEADERS, ["=1+1", 123, "Company A", "Staff", "Marketing"]], "formula"),
    ([HEADERS, ["User", True, "Company A", "Staff", "Marketing"]], "boolean"),
    ([HEADERS, ["User", 123, "Company A", "Staff", "Marketing", "secret"]], "di luar"),
])
def test_parser_validation(payload, message):
    with pytest.raises(ValueError, match=message):
        read_user_import("bad.xlsx", xlsx(payload))


@pytest.mark.parametrize("extension", [".xls", ".xlsx", ".csv", ".xlsm"])
def test_parser_rejects_bad_files(extension):
    with pytest.raises(ValueError):
        read_user_import("bad" + extension, b"not an Excel file")
    with pytest.raises(ValueError):
        read_user_import("bad" + extension, b"")


def test_parser_size_rows_zip_bomb_and_ambiguous_sheets():
    with pytest.raises(ValueError, match="5 MB"):
        read_user_import("large.xlsx", b"x" * (MAX_USER_IMPORT_BYTES + 1))
    with pytest.raises(ValueError, match="500"):
        read_user_import("large.xlsx", xlsx([HEADERS] + [["User", 123, "Company A", "Staff", "Marketing"]] * 501))
    bomb = BytesIO()
    with ZipFile(bomb, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * (21 * 1024 * 1024))
    with pytest.raises(ValueError, match="diekstrak"):
        read_user_import("bomb.xlsx", bomb.getvalue())
    book = openpyxl.Workbook()
    book.create_sheet("Another")
    output = BytesIO()
    book.save(output)
    with pytest.raises(ValueError, match="multi-sheet"):
        read_user_import("multi.xlsx", output.getvalue())


def snapshot(database):
    with sqlite3.connect(database.path) as connection:
        return {table: connection.execute(f"SELECT * FROM {table} ORDER BY 1").fetchall()
                for table in ("users", "user_company_memberships", "user_sessions", "messages", "module_access")}


def test_existing_active_and_inactive_ids_are_entirely_untouched(database):
    for ident, active in [(100, True), (200, False)]:
        database.create_user_with_membership(ident, "Original", "company-a", "Owner", "Management",
                                             "gm", "executive", "Keep this", "admin", active)
    database.set_active_company(100, "company-a")
    database.add_message(100, "user", "Keep history", "company-a")
    before = snapshot(database)
    # Invalid Excel attributes on an existing ID must not block or change it.
    report = run_import(database, [row(2, 100, "Missing Company", name=""), row(3, 200, "Company B")], active=True)
    assert report.created_users == 0 and report.skipped_rows == (2, 3)
    assert snapshot(database) == before
    assert database.get_membership_admin(200, "company-b") is None


def test_new_multi_company_duplicate_and_repeat_import(database):
    rows = [row(company=" company a "), row(3, company="Company B"), row(4, company="company-a")]
    report = run_import(database, rows)
    assert report.created_users == 1 and report.created_memberships == 2
    assert report.duplicate_rows == (4,)
    assert database.get_user_admin(123456789).active is False
    assert database.get_membership_admin(123456789, "company-a").is_default is True
    assert database.get_membership_admin(123456789, "company-b").is_default is False
    assert database.list_accessible_modules(123456789, "company-a") == []
    before = snapshot(database)
    again = run_import(database, rows, active=True)
    assert again.created_users == 0 and again.skipped_rows == (2, 3, 4)
    assert snapshot(database) == before


def test_id_precision_and_admin_defaults_validation(database):
    with pytest.raises(UserImportValidationError, match="Import dibatalkan"):
        run_import(database, [row(telegram_id=9007199254740990)])
    assert run_import(database, [row(telegram_id="9007199254740990")]).created_users == 1
    for role, profile in [("owner", "staff"), ("staff", "unknown")]:
        with pytest.raises(ValueError, match="tidak valid"):
            database.import_new_users([row()], role_level=role, communication_profile=profile, active=False, actor="admin")
    assert database.get_user_admin(123456789) is None


def test_mixed_existing_and_new_does_not_modify_existing(database):
    run_import(database, [row(telegram_id=100)])
    original = database.get_user_admin(100)
    original_membership = database.get_membership_admin(100, "company-a")
    report = run_import(database, [row(telegram_id=100, name="Changed", company="Company B"), row(3, 200)], active=True)
    assert report.created_users == 1 and report.skipped_rows == (2,)
    assert database.get_user_admin(100) == original
    assert database.get_membership_admin(100, "company-a") == original_membership
    assert database.get_membership_admin(100, "company-b") is None
    assert database.get_user(200) is not None


@pytest.mark.parametrize("invalid", [
    row(3, 234, "Missing"), row(3, "@username"), row(3, 1.5), row(3, 0),
    row(3, 2**53), row(3, float("nan")), row(3, True),
    row(3, 234, name="A"), row(3, 234, division=""),
    row(3, name="Different"), row(3, job_title="Different"),
])
def test_invalid_new_rows_roll_back_entire_batch(database, invalid):
    before = snapshot(database)
    with pytest.raises(UserImportValidationError) as exc:
        run_import(database, [row(), invalid])
    assert "Baris 3" in exc.value.errors[0]
    assert snapshot(database) == before


def test_inactive_or_ambiguous_company_rejected(database):
    database.set_company_active("company-b", False, "admin")
    with pytest.raises(UserImportValidationError):
        run_import(database, [row(company="Company B")])
    database.create_company("", "Company A", "admin")
    with pytest.raises(UserImportValidationError) as exc:
        run_import(database, [row()])
    assert "ambigu" in str(exc.value.errors)
    assert run_import(database, [row(company="company-a")]).created_users == 1


def test_concurrent_imports_do_not_overwrite(database):
    with ThreadPoolExecutor(max_workers=2) as executor:
        reports = list(executor.map(lambda _: run_import(database, [row()]), range(2)))
    assert sorted(report.created_users for report in reports) == [0, 1]


def test_transaction_rolls_back_on_write_failure(database):
    with sqlite3.connect(database.path) as connection:
        connection.execute("""CREATE TRIGGER reject_membership BEFORE INSERT ON user_company_memberships
                              BEGIN SELECT RAISE(ABORT, 'test failure'); END""")
        audits_before = connection.execute("SELECT COUNT(*) FROM admin_audit_events").fetchone()[0]
    with pytest.raises(sqlite3.IntegrityError):
        run_import(database, [row()])
    assert database.get_user_admin(123456789) is None
    with sqlite3.connect(database.path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM admin_audit_events").fetchone()[0] == audits_before


@pytest.mark.parametrize("writer, extension", [(xlsx, ".xlsx"), (xls, ".xls")])
def test_admin_import_end_to_end_and_authorization(database, tmp_path, writer, extension):
    settings = _test_settings(tmp_path, tmp_path / "users.json", tmp_path / "companies.json", database_path=database.path)
    content = writer([HEADERS, ["Ajeng", "106545875", "Company A", "Manager", "Marketing"]])
    with TestClient(create_admin_app(settings, database)) as client:
        assert client.get("/admin/users/import", follow_redirects=False).status_code == 303
        assert client.post("/admin/users/import", files={"file": ("users" + extension, content)}, follow_redirects=False).status_code == 303
        client.post("/admin/login", data={"username": "admin", "password": "strong-password"})
        assert 'href="/admin/users/import"' in client.get("/admin/users").text
        page = client.get("/admin/users/import")
        csrf = re.search(r'name="csrf_token" value="([a-f0-9]+)"', page.text).group(1)
        assert client.post("/admin/users/import", files={"file": ("users" + extension, content)}).status_code == 403
        assert database.get_user_admin(106545875) is None
        response = client.post("/admin/users/import", data={"csrf_token": csrf, "role_level": "manager", "communication_profile": "manager", "active": "1"}, files={"file": ("users" + extension, content)})
        assert response.status_code == 200
        assert "1 user baru" in response.text
        assert database.get_user(106545875).name == "Ajeng"
        membership = database.get_membership_admin(106545875, "company-a")
        assert membership.role_level == "manager" and membership.communication_profile == "manager"
        before = snapshot(database)
        retry = client.post("/admin/users/import", data={"csrf_token": csrf}, files={"file": ("users" + extension, content)})
        assert retry.status_code == 200 and "0 user baru" in retry.text
        assert snapshot(database) == before
        assert "Import user baru" in client.get("/admin/activity").text
        assert client.post("/admin/users/import", data={"csrf_token": csrf}).status_code == 400
        assert client.post("/admin/users/import", data={"csrf_token": csrf}, files={"file": ("bad.xlsx", b"bad")}).status_code == 400
        bad = client.post("/admin/users/import", data={"csrf_token": csrf}, files={"file": ("bad.xlsx", xlsx([HEADERS, ["<script>alert(1)</script>", 321, "Missing", "Staff", "Marketing"]]))})
        assert bad.status_code == 400 and "Baris 2" in bad.text
        assert "<script>alert(1)</script>" not in bad.text
        assert database.get_user_admin(321) is None
        assert client.post("/admin/users/import", data={"csrf_token": csrf}, files={"file": ("big.xlsx", b"x" * (MAX_USER_IMPORT_BYTES + 100_000))}).status_code == 413


def test_request_limit_works_without_content_length():
    async def check():
        async def receive():
            return {"type": "http.request", "body": b"x" * (MAX_USER_IMPORT_BYTES + 100_000), "more_body": True}
        request = Request({"type": "http", "method": "POST", "path": "/admin/users/import", "headers": []}, receive)
        with pytest.raises(HTTPException) as exc:
            await _bounded_import_request(request)
        assert exc.value.status_code == 413
    asyncio.run(check())
