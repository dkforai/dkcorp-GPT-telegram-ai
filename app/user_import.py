"""Bounded Excel decoding for the five-column user template; no database writes."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

MAX_USER_IMPORT_BYTES = 5 * 1024 * 1024
MAX_USER_IMPORT_ROWS = 500
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_COLUMNS = 10
HEADERS = ("Nama", "Telegram ID", "Perusahaan", "Jabatan", "Divisi")


@dataclass(frozen=True)
class UserImportRow:
    row_number: int
    name: object
    telegram_id: object
    company: object
    job_title: object
    division: object


@dataclass(frozen=True)
class UserImportReport:
    created_users: int
    created_memberships: int
    skipped_rows: tuple[int, ...]
    duplicate_rows: tuple[int, ...]


class UserImportValidationError(ValueError):
    def __init__(self, errors: list[str]):
        super().__init__("Import dibatalkan. Belum ada data yang disimpan.")
        self.errors = errors


def _header(value: object) -> str:
    return " ".join(str(value or "").strip().casefold().replace("_", " ").split())


def _has_value(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def read_user_import(filename: str, data: bytes) -> list[UserImportRow]:
    extension = Path(filename.replace("\\", "/")).suffix.casefold()
    if extension not in {".xls", ".xlsx"}:
        raise ValueError("Upload file Excel .xls atau .xlsx, bukan format lainnya.")
    if not data:
        raise ValueError("File Excel kosong.")
    if len(data) > MAX_USER_IMPORT_BYTES:
        raise ValueError("Ukuran file Excel maksimal 5 MB.")
    try:
        grid = _read_xlsx(data) if extension == ".xlsx" else _read_xls(data)
    except ValueError:
        raise
    except Exception as exc:
        # Parser errors may contain workbook contents. Do not expose/log them.
        raise ValueError(
            "File Excel rusak, dilindungi password, atau isinya tidak sesuai ekstensi. "
            "Simpan ulang sebagai .xls/.xlsx tanpa password."
        ) from exc
    if not grid:
        raise ValueError("File Excel tidak memiliki header.")
    header = [_header(value) for value in grid[0]]
    while header and not header[-1]:
        header.pop()
    required = [_header(value) for value in HEADERS]
    if len(header) != len(required) or set(header) != set(required):
        raise ValueError(
            "Baris 1 harus berisi tepat lima kolom: Nama, Telegram ID, "
            "Perusahaan, Jabatan, Divisi. Gunakan template sederhana."
        )
    positions = [header.index(label) for label in required]
    rows = []
    for number, values in enumerate(grid[1:], start=2):
        if not any(_has_value(value) for value in values):
            continue
        if any(_has_value(value) for value in values[len(header):]):
            raise ValueError(f"Baris {number}: ada data di luar lima kolom template.")
        padded = values + [None] * (len(header) - len(values))
        rows.append(UserImportRow(number, *(padded[i] for i in positions)))
    if not rows:
        raise ValueError("Belum ada data user. Isi mulai baris 2.")
    return rows


def _read_xlsx(data: bytes) -> list[list[object]]:
    import openpyxl

    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
            raise ValueError("Isi Excel terlalu besar setelah diekstrak.")
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError("Excel dengan password tidak didukung.")
    book = openpyxl.load_workbook(
        BytesIO(data), read_only=True, data_only=False, keep_links=False,
    )
    try:
        if not book.worksheets:
            raise ValueError("Excel tidak memiliki worksheet.")
        if "Data_User" in book.sheetnames:
            sheet = book["Data_User"]
        elif len(book.worksheets) == 1:
            sheet = book.worksheets[0]
        else:
            raise ValueError("Untuk file multi-sheet, beri nama tab data user Data_User.")
        # Do not trust worksheet dimensions supplied by an uploaded document.
        sheet.reset_dimensions()
        grid = []
        for number, cells in enumerate(sheet.iter_rows(), start=1):
            if number > MAX_USER_IMPORT_ROWS + 1:
                raise ValueError("Excel maksimal 500 baris data (baris 2–501).")
            if len(cells) > MAX_COLUMNS:
                raise ValueError("Terlalu banyak kolom. Gunakan template lima kolom.")
            if any(cell.data_type in {"f", "e"} for cell in cells):
                raise ValueError(f"Baris {number}: gunakan nilai biasa, bukan formula/error Excel.")
            values = [cell.value for cell in cells]
            _check_scalars(values, number)
            grid.append(values)
        return grid
    finally:
        book.close()


def _read_xls(data: bytes) -> list[list[object]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data, on_demand=True, ragged_rows=True, logfile=StringIO())
    try:
        names = book.sheet_names()
        if not names:
            raise ValueError("Excel tidak memiliki worksheet.")
        if "Data_User" in names:
            sheet = book.sheet_by_name("Data_User")
        elif len(names) == 1:
            sheet = book.sheet_by_index(0)
        else:
            raise ValueError("Untuk file multi-sheet, beri nama tab data user Data_User.")
        if sheet.nrows > MAX_USER_IMPORT_ROWS + 1 or sheet.ncols > MAX_COLUMNS:
            raise ValueError("Excel maksimal 500 baris data dan memakai template lima kolom.")
        grid = []
        for index in range(sheet.nrows):
            cells = sheet.row(index)
            if any(c.ctype in {xlrd.XL_CELL_DATE, xlrd.XL_CELL_BOOLEAN, xlrd.XL_CELL_ERROR} for c in cells):
                raise ValueError(f"Baris {index + 1}: gunakan teks/angka, bukan tanggal, boolean, atau error Excel.")
            values = [cell.value for cell in cells]
            _check_scalars(values, index + 1)
            grid.append(values)
        return grid
    finally:
        book.release_resources()


def _check_scalars(values: list[object], number: int) -> None:
    if any(
        value is not None and (
            isinstance(value, bool) or not isinstance(value, (str, int, float))
            or isinstance(value, str) and len(value) > 1000
        ) for value in values
    ):
        raise ValueError(f"Baris {number}: isi sel harus teks/angka singkat, bukan tanggal atau boolean.")
