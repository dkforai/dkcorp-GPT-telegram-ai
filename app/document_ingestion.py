from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_EXTRACTED_CHARS = 100_000
MAX_PDF_PAGES = 500
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    media_type: str
    size_bytes: int
    sha256: str
    text: str


def extract_uploaded_document(
    filename: object, media_type: object, data: bytes
) -> ExtractedDocument:
    safe_name = Path(str(filename or "").replace("\\", "/")).name.strip()
    if not safe_name:
        raise ValueError("Nama file upload tidak valid")
    extension = Path(safe_name).suffix.casefold()
    if extension == ".doc":
        raise ValueError(
            "Format Word lama .doc belum didukung. Simpan ulang sebagai .docx lalu upload kembali"
        )
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Format file harus PDF, DOCX, TXT, atau Markdown")
    if not data:
        raise ValueError("File upload kosong")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("Ukuran file maksimal 10 MB")

    normalized_media_type = str(media_type or "application/octet-stream").strip()
    if extension == ".pdf":
        text = _extract_pdf(data)
    elif extension == ".docx":
        text = _extract_docx(data)
    else:
        text = _extract_plain_text(data)
    text = _normalize_text(text)
    if not text:
        if extension == ".pdf":
            raise ValueError(
                "PDF tidak mengandung teks yang dapat diekstrak. PDF hasil scan memerlukan OCR"
            )
        raise ValueError("File tidak mengandung teks yang dapat digunakan")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise ValueError(
            "Hasil ekstraksi melebihi 100.000 karakter. Pecah dokumen menjadi beberapa file"
        )
    return ExtractedDocument(
        filename=safe_name[:255],
        media_type=normalized_media_type[:150],
        size_bytes=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        text=text,
    )


def _extract_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise ValueError("Isi file tidak cocok dengan format PDF")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ValueError("PDF yang dilindungi password belum didukung")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF maksimal 500 halaman")
        pages: list[str] = []
        extracted_chars = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)
            extracted_chars += len(page_text)
            if extracted_chars > MAX_EXTRACTED_CHARS:
                break
        return "\n\n".join(pages)
    except ValueError:
        raise
    except (PdfReadError, OSError) as exc:
        raise ValueError("PDF rusak atau tidak dapat dibaca") from exc


def _extract_docx(data: bytes) -> str:
    if not data.startswith(b"PK"):
        raise ValueError("Isi file tidak cocok dengan format DOCX")
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                raise ValueError("File bukan dokumen Word DOCX yang valid")
            uncompressed_size = sum(item.file_size for item in archive.infolist())
            if uncompressed_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("Isi DOCX terlalu besar setelah diekstrak")
        document = Document(BytesIO(data))
    except ValueError:
        raise
    except (BadZipFile, KeyError, OSError, PackageNotFoundError) as exc:
        raise ValueError("DOCX rusak atau tidak dapat dibaca") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_plain_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("File teks harus memakai encoding UTF-8") from exc


def _normalize_text(text: str) -> str:
    normalized_newlines = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in normalized_newlines.split("\n")]
    normalized: list[str] = []
    blank = False
    for line in lines:
        if line.strip():
            normalized.append(line.strip())
            blank = False
        elif normalized and not blank:
            normalized.append("")
            blank = True
    return "\n".join(normalized).strip()
